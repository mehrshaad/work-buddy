"""Checks for the worklog engine. Run: python3 tests/test_worklog.py

Lives in the repo on purpose — an earlier version sat in a session scratchpad and was
deleted with it. Uses the real config and the real data on this machine, so some checks
are skipped when a source is unavailable rather than failing.
"""
import json, sys, tempfile
from datetime import date as D, datetime, timedelta
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "bin"))
import worklog as W

CFG = json.loads((Path.home() / ".worklog" / "config.json").read_text())
ok, fail, skip = [], [], []
TMP = Path(tempfile.mkdtemp(prefix="worklog-tests-"))


def check(name, cond, detail=""):
    (ok if cond else fail).append(f"{name}{' — ' + detail if detail else ''}")


def stub(name, body):
    f = TMP / name
    f.write_text(body)
    f.chmod(0o755)
    return str(f)


# ---------------------------------------------------------------- claude code --
cc = W.collect_claude_code(CFG, D(2026, 7, 30))
if not cc.get("session_count"):
    skip.append("claude_code: no sessions on the sample day")
else:
    ss = cc["sessions"]
    check("cc sessions have a duration and turn count",
          all(s.get("duration") and "assistant_turns" in s for s in ss))
    check("cc at least one session carries a title",
          any(s.get("title") for s in ss), str([s.get("title") for s in ss])[:120])
    check("cc a session records the git branch",
          any(s.get("branches") for s in ss))
    check("cc command intents are captured",
          any(s.get("did") for s in ss))
    check("cc changed files are captured",
          any(s.get("files_changed") for s in ss))
    check("cc file paths are never absolute",
          all(not f.startswith("/Users/") for s in ss for f in s.get("files_changed", [])),
          str([f for s in ss for f in s.get("files_changed", []) if f.startswith("/Users/")])[:100])
    check("cc did-list respects its cap",
          all(len(s.get("did", [])) <= CFG["claude_code"]["max_commands_per_session"] for s in ss))
    check("cc files respect their cap",
          all(len(s.get("files_changed", [])) <= CFG["claude_code"]["max_files_per_session"] for s in ss))
    check("cc no prompt is harness noise",
          not any(W._PROMPT_NOISE.match(p) for s in ss for p in s.get("prompts", [])),
          str([p[:40] for s in ss for p in s.get("prompts", []) if W._PROMPT_NOISE.match(p)])[:120])
    check("cc left_off is a real prompt, not injected text",
          all(not W._PROMPT_NOISE.match(s["left_off"]) for s in ss if s.get("left_off")))
    # `prompts` is capped at the first N, which show intent; `left_off` is the last
    # prompt of the day, so on a long session it is deliberately not in that list
    _cap = CFG["claude_code"]["max_prompts_per_session"]
    check("cc left_off equals the last prompt when nothing was truncated",
          all(s["left_off"] == s["prompts"][-1] for s in ss
              if s.get("left_off") and 0 < len(s["prompts"]) < _cap))
    check("cc left_off is set whenever the session had prompts",
          all(bool(s.get("left_off")) == bool(s.get("prompts")) for s in ss))
    check("cc no two sessions share start+end+turns (forks collapsed)",
          len({(s["start"], s["end"], s["assistant_turns"], s["project"]) for s in ss}) == len(ss))
    check("cc self-summarizer sessions are excluded",
          not any(W.is_self_prompt(p) for s in ss for p in s.get("prompts", [])))

check("cc noise filter catches interruptions",
      not W._is_real_prompt("[Request interrupted by user]"))
check("cc noise filter catches injected worktree headers",
      not W._is_real_prompt("Repo worktree: /Users/x/y (branch z)"))
check("cc noise filter catches skill preamble",
      not W._is_real_prompt("Base directory for this skill: /tmp/x"))
check("cc noise filter keeps ordinary prompts",
      W._is_real_prompt("fix the retry logic and push"))
check("cc _short_path strips the home prefix",
      not W._short_path("/Users/x/Documents/proj/a/B.java").startswith("/"),
      W._short_path("/Users/x/Documents/proj/a/B.java"))
check("cc test files are recognised",
      bool(W._TEST_HINT.search("backend/src/test/java/FooTest.java"))
      and not W._TEST_HINT.search("backend/src/main/java/Foo.java"))
check("cc _top_n keeps first-seen order when asked",
      W._top_n(["b", "a", "b", "c"], 2, keep_order=True) == ["b", "a"])
check("cc _top_n ranks by frequency otherwise",
      W._top_n(["b", "a", "b", "c"], 1) == ["b"])

_forked = [
    {"start": "s", "end": "e", "assistant_turns": 9, "project": "p", "title": "A", "ai_title": ""},
    {"start": "s", "end": "e", "assistant_turns": 9, "project": "p", "title": "B", "ai_title": ""},
    {"start": "s2", "end": "e2", "assistant_turns": 3, "project": "p", "title": "C", "ai_title": ""},
]
_kept, _dropped = W._dedupe_forked_sessions(CFG, _forked)
check("fork dedupe collapses identical transcripts", len(_kept) == 2 and _dropped == 1)
check("fork dedupe keeps the alternate title",
      "B" in (_kept[0].get("also_titled") or []), str(_kept[0].get("also_titled")))

# ------------------------------------------------------------------ summarize --
_dig = {"date": "2026-07-30", "weekday": "Thursday",
        "window": {"start": "2026-07-30T08:00:00", "end": "2026-07-30T17:00:00"},
        "generated_at": "2026-07-30T17:00:00",
        "sources": {"activity": {"available": True, "apps": [
                        {"app": "Code", "minutes": 30, "duration": "30m", "titles": []}],
                        "active_duration": "30m", "idle_duration": "0m"},
                    "git": {"available": True, "total_commits": 0},
                    "shell": {"available": True, "count": 0},
                    "calendar": {"available": False, "reason": "x"},
                    "meetings_attended": {"available": True, "count": 0},
                    "unjoined_meetings": {"available": False, "reason": "x"},
                    "mis_board": {"available": True, "count": 0},
                    "claude_code": {"available": True, "session_count": 0},
                    "cloud": {"available": True, "count": 0},
                    "claude_export": {"available": False, "count": 0}}}


def run_summarize(binary, attempts):
    c = json.loads(json.dumps(CFG))
    c["state_dir"] = str(TMP / "state")
    (TMP / "state" / "logs").mkdir(parents=True, exist_ok=True)
    for f in (TMP / "state" / "logs").glob("*.log"):
        f.unlink()
    c["summarizer"].update({"claude_binary": binary, "attempts": attempts,
                            "retry_delay_seconds": 0})
    md, method = W.summarize(c, _dig)
    log = "".join(f.read_text() for f in (TMP / "state" / "logs").glob("*.log"))
    return md, method, log


_fail = stub("fail.sh", '#!/bin/sh\necho "usage limit reached" >&2\nexit 1\n')
_md, _m, _log = run_summarize(_fail, 3)
check("summarize retries three times", _log.count("claude CLI failed") == 3,
      str(_log.count("claude CLI failed")))
check("summarize falls back only after the last attempt", _m == "fallback")
check("summarize logs the actual reason", "usage limit reached" in _log)
_, _m1, _log1 = run_summarize(_fail, 1)
check("summarize attempts=1 makes one attempt", _log1.count("claude CLI failed") == 1)

_flag = TMP / "tried"
_flap = stub("flap.sh", f'#!/bin/sh\nif [ -f "{_flag}" ]; then printf "**Work update**\\n\\n- x\\n"; '
                        f'exit 0; fi\ntouch "{_flag}"; echo "connection reset" >&2; exit 1\n')
_md2, _m2, _log2 = run_summarize(_flap, 3)
check("summarize recovers on a later attempt", _m2 == "claude_cli", _m2)
check("summarize emits no fallback text after recovering",
      "Generated without summarization" not in _md2)
_sout = stub("sout.sh", '#!/bin/sh\necho "Error: not logged in"\nexit 1\n')
_, _, _log3 = run_summarize(_sout, 2)
check("summarize surfaces errors printed on stdout", "not logged in" in _log3)

# -------------------------------------------------------------------- health ---
_hp = TMP / "h.md"
_body = ("<!-- worklog claude_cli | generated x -->\n\n**Work update**\n\n---\n"
         "**MIS Board Tasks**\n\n| TBD | a | Sam | Medium | Done | ~1h | d |\n")
_tail = "\n---\n**After-hours addendum** _(x)_\n\n**Shell**\n- `ls`\n"
_hp.write_text(_body)
check("health accepts a good report", W.report_health(_hp)[0] == "ok", str(W.report_health(_hp)))
_hp.write_text(_body.replace("claude_cli", "fallback"))
check("health flags a fallback report", W.report_health(_hp) == ("degraded", "written via fallback"))
_hp.write_text(_body.replace("**MIS Board Tasks**", "**Something Else**"))
check("health flags a missing tracker table", W.report_health(_hp)[0] == "degraded")
# the section was renamed with the skill; both spellings must stay valid or every
# report written before the rename would suddenly read as degraded
_hp.write_text(_body.replace("**MIS Board Tasks**", W.SOR_MARKER))
check("health accepts the new SOR Status Tracker marker", W.report_health(_hp)[0] == "ok")
check("health still accepts the legacy MIS marker",
      W.MIS_MARKER in W.TRACKER_MARKERS and W.SOR_MARKER in W.TRACKER_MARKERS)
check("rows parse under the new marker",
      len(W._mis_rows(_body.replace("**MIS Board Tasks**", W.SOR_MARKER))) == 1)
check("rows parse under the legacy marker", len(W._mis_rows(_body)) == 1)
_newtbl = ("<!-- worklog claude_cli | x -->\n\n**Work update**\n\n---\n" + W.SOR_MARKER +
           "\n\n| Agent | Action Item | Owner(s) | Priority | Status | Time Spent | Description |\n"
           "|---|---|---|---|---|---|---|\n"
           "| Project A | a thing | Sam | High | Done | ~1h | did it |\n")
check("the new Agent-led header is skipped, not parsed as a row",
      W._mis_rows(_newtbl) == ["| Project A | a thing | Sam | High | Done | ~1h | did it |"],
      str(W._mis_rows(_newtbl)))

# the Agent dropdown must reach the prompt byte-for-byte: a "corrected" spelling of
# a corrected string does not match the sheet's dropdown and the card lands nowhere
_agents = (CFG.get("tracker") or {}).get("agents") or []
if not _agents:
    skip.append("tracker.agents not configured")
else:
    _al = "\n".join(f"    {a}" for a in _agents)
    _tk = CFG.get("tracker") or {}
    _p = W.SUMMARY_PROMPT.format(
        date_h="x", next_up="", agent_list=_al,
        agent_hints=_tk.get("agent_hints", ""),
        owner=_tk.get("default_owner", "me"),
        section=_tk.get("section_title", W.DEFAULT_SECTION), payload="{}")
    check("every configured project reaches the prompt verbatim",
          all(a in _p for a in _agents), str(len(_agents)))
    _eng = Path(W.__file__).read_text()
    check("the project list is not hardcoded in the engine",
          not any(a in _eng for a in _agents))
    check("the owner name is not hardcoded in the engine",
          _tk.get("default_owner", "\x00") not in _eng)
    check("the project hints are not hardcoded in the engine",
          not _tk.get("agent_hints") or _tk["agent_hints"][:40] not in _eng)
    _hdr = "| Agent | Action Item | Owner(s) | Priority | Status | Time Spent | Description |"
    check("the daily table leads with Agent and has no card number",
          _hdr in W.SUMMARY_PROMPT and "S. NO" not in W.SUMMARY_PROMPT)
    check("the section title comes from config, not the engine",
          "{section}" in W.SUMMARY_PROMPT and "SOR" not in W.SUMMARY_PROMPT)
    check("the weekly table leads with Agent and has no card number",
          _hdr in W.WEEKLY_PROMPT and "S. NO" not in W.WEEKLY_PROMPT)
    check("the fallback renderer emits the same header",
          _hdr in Path(W.__file__).read_text())
    check("no placeholder is used when no project fits",
          "NOT write a placeholder" in W.SUMMARY_PROMPT)
    check("the weekly prompt tolerates six, seven or eight cells",
          "six, seven or eight cells" in W.WEEKLY_PROMPT)

# ------------------------------------------------------------- excluded apps ---
for _a in ("FaceTime", "Steam", "Spotify", "zoom.us", "Zoom", "Telegram", "telegram", "ZOOM.US"):
    check(f"{_a} is excluded", W.is_excluded_app(CFG["activity"], _a))
for _a in ("Code", "MSTeams", "Google Chrome"):
    check(f"{_a} is not excluded", not W.is_excluded_app(CFG["activity"], _a))
check("Zoom is not a meeting source",
      not any("zoom" in a.lower() for a in CFG["activity"]["meeting_apps"]))

# vendor rule: every com.apple.* app is excluded unless explicitly allowed, so an
# Apple app never opened before is covered without maintaining a list of names
_A = CFG["activity"]
for _keep in ("Notes", "Reminders", "System Settings", "Calendar", "Activity Monitor"):
    _b = W.bundle_for_app(_keep)
    check(f"{_keep} is kept", _b and not W.is_excluded_app(_A, _keep), str(_b))
for _drop in ("Finder", "Safari", "Mail", "Photos", "Music", "Maps", "Preview",
              "TextEdit", "Calculator", "App Store", "Freeform", "Stickies"):
    _b = W.bundle_for_app(_drop)
    check(f"{_drop} is excluded", _b and W.is_excluded_app(_A, _drop), str(_b))
for _other in ("Google Chrome", "Microsoft Outlook", "Postman"):
    check(f"{_other} is unaffected by the vendor rule", not W.is_excluded_app(_A, _other))
check("an unknown com.apple bundle is excluded on sight",
      W.is_excluded_app(_A, "Some Future App", "com.apple.SomethingNew"))
check("an allowed bundle passed directly is kept",
      not W.is_excluded_app(_A, "Whatever", "com.apple.Notes"))
check("vendor matching ignores bundle-id case",
      W.is_excluded_app(_A, "x", "COM.APPLE.SAFARI")
      and not W.is_excluded_app(_A, "x", "COM.APPLE.NOTES"))
check("placeholder app names are never excluded",
      not W.is_excluded_app(_A, "(idle)") and not W.is_excluded_app(_A, "(unknown)"))
check("bundle_for_app ignores placeholders", W.bundle_for_app("(idle)") is None)
check("no vendor prefixes configured means the rule is off",
      not W.is_excluded_app({"exclude_apps": []}, "Safari"))
check("no banned app name survives in the raw record",
      not any(W.is_excluded_app(CFG["activity"], json.loads(l).get("app", ""))
              for f in (Path.home() / ".worklog" / "raw").glob("*/activity.jsonl")
              for l in f.read_text(errors="replace").splitlines() if l.strip()))

_es = TMP / "excl"; (_es / "raw" / "2026-07-28").mkdir(parents=True, exist_ok=True)
_cfge = json.loads(json.dumps(CFG)); _cfge["state_dir"] = str(_es)
_rows = ([{"ts": f"2026-07-28T09:{i:02d}:00", "app": "Code", "title": "", "idle": 0} for i in range(10)]
         + [{"ts": f"2026-07-28T09:{i:02d}:00", "app": "Telegram", "title": "chat", "idle": 0}
            for i in range(10, 40)])
(_es / "raw" / "2026-07-28" / "activity.jsonl").write_text(
    "".join(json.dumps(r) + "\n" for r in _rows))
_act = W.collect_activity(_cfge, D(2026, 7, 28))
check("excluded app never appears as an app", [a["app"] for a in _act["apps"]] == ["Code"])
check("excluded time is not counted as idle", _act["idle_duration"] == "0m", _act["idle_duration"])

# -------------------------------------------------------- meetings + weekly ----
MEET = "Meeting in X | Acme | a@b.com | Microsoft Teams"
CHAT = "Chat | Acme | a@b.com | Microsoft Teams"
_ms = TMP / "meet"; (_ms / "raw" / "2026-07-29").mkdir(parents=True, exist_ok=True)
_cfgm = json.loads(json.dumps(CFG)); _cfgm["state_dir"] = str(_ms)
_pres = [{"ts": f"2026-07-29T09:{i:02d}:00", "app": "Code", "title": "", "idle": 900,
          "meeting_windows": [["MSTeams", MEET], ["MSTeams", CHAT]]} for i in range(31)]
(_ms / "raw" / "2026-07-29" / "activity.jsonl").write_text(
    "".join(json.dumps(r) + "\n" for r in _pres))
_att = W.collect_meetings_attended(_cfgm, D(2026, 7, 29))
check("meeting credited while its window is open, even when idle and unfocused",
      _att["count"] == 1 and _att["meetings"][0]["minutes"] >= 29,
      str(_att["meetings"][0]["minutes"] if _att["count"] else None))
check("meeting basis is window presence", _att["basis"] == "window presence")
check("chat window is not a meeting", _att["count"] == 1)
check("focused_minutes is no longer reported",
      "focused_minutes" not in _att["meetings"][0])
_twice = ([{"ts": f"2026-07-29T09:{i:02d}:00", "app": "Code", "title": "", "idle": 0,
            "meeting_windows": [["MSTeams", MEET]]} for i in range(10)]
          + [{"ts": f"2026-07-29T16:{i:02d}:00", "app": "Code", "title": "", "idle": 0,
              "meeting_windows": [["MSTeams", MEET]]} for i in range(10)])
(_ms / "raw" / "2026-07-29" / "activity.jsonl").write_text(
    "".join(json.dumps(r) + "\n" for r in _twice))
check("a meeting held twice is two occurrences",
      W.collect_meetings_attended(_cfgm, D(2026, 7, 29))["count"] == 2)

check("week runs Saturday to Friday",
      W.week_sat_to_fri(D(2026, 7, 31)) == (D(2026, 7, 25), D(2026, 7, 31)))
check("any weekday maps to that week's Friday",
      W.week_sat_to_fri(D(2026, 7, 27)) == (D(2026, 7, 25), D(2026, 7, 31)))
_mp = W.expand((CFG.get("mis_board") or {}).get("path", "~/.worklog/mis-board.xlsx"))
if not _mp.is_file():
    skip.append("mis_board: workbook not present")
else:
    _rows2, _org = W._mis_rows_for(CFG)
    check("mis board rows are readable", _rows2 and len(_rows2) > 10, f"{_org}")
    check("mis board rows are the configured owner's",
          all(CFG["mis_board"]["owner_match"] in r["owners"] for r in _rows2))
    check("mis board never exposes hours",
          not any("hours" in k.lower() for r in _rows2 for k in r))
    _mo = TMP / "misweek"; _mo.mkdir(exist_ok=True)
    _cfgw = json.loads(json.dumps(CFG)); _cfgw["output_dir"] = str(_mo)
    _note = W.write_mis_week(_cfgw, D(2026, 7, 24))
    check("weekly excel is named by ISO week", "mis-week-2026-W30.xlsx" in (_note or ""), str(_note))
    import openpyxl
    _out = openpyxl.load_workbook(_mo / "mis-week-2026-W30.xlsx")
    _src = openpyxl.load_workbook(_mp)
    check("weekly excel preserves the sheets", _out.sheetnames == _src.sheetnames)
    check("weekly excel preserves the columns",
          [c.value for c in _out["Work items"][1]] == [c.value for c in _src["Work items"][1]])

# --------------------------------------------------------------------- agents --
for _lbl, _cmd in (("summarize", "report"), ("catchup", "repair"), ("tracker", "sample")):
    _pl = Path.home() / f"Library/LaunchAgents/com.workbuddy.{_lbl}.plist"
    check(f"{_lbl} agent is installed", _pl.is_file())
    if _pl.is_file():
        check(f"{_lbl} agent runs `{_cmd}`", f"<string>{_cmd}</string>" in _pl.read_text())
_cu = Path.home() / "Library/LaunchAgents/com.workbuddy.catchup.plist"
if _cu.is_file():
    import plistlib as _plib
    _d = _plib.load(open(_cu, "rb"))
    check("catchup runs at 10:00 on five weekdays",
          sorted((x["Weekday"], x["Hour"]) for x in _d["StartCalendarInterval"])
          == [(i, 10) for i in range(1, 6)],
          str(_d["StartCalendarInterval"]))
    check("no stale com.worklog agent is still loaded",
          not list((Path.home()/"Library/LaunchAgents").glob("com.worklog.*.plist")))

# ---------------------------------------------------------------------- pause --
_ps = TMP / "pause"; (_ps / "logs").mkdir(parents=True, exist_ok=True)
_cfgp = json.loads(json.dumps(CFG)); _cfgp["state_dir"] = str(_ps)
_N = datetime(2026, 8, 12, 10, 0, 0)


def _win(a, b):
    """Write one completed pause window straight into the history."""
    (_ps / "pauses.jsonl").write_text(json.dumps(
        {"started": a.isoformat(), "ended": b.isoformat(),
         "reason": "personal", "ended_by": "resume"}) + "\n")


check("no pause means no active pause", W.active_pause(_cfgp, _N) is None)
check("rest-of-day refuses outside the work window rather than pausing until tomorrow",
      "rest_of_day" in Path(W.__file__).read_text()
      and "nothing to pause" in Path(W.__file__).read_text())
check("no pause means no windows", W.pause_windows(_cfgp, D(2026, 8, 12)) == [])
check("status line reads tracking when clear", W.pause_status_line(_cfgp, _N) == "tracking")

_r = W.start_pause(_cfgp, minutes=45)
_a = W.active_pause(_cfgp, _N)
check("a timed pause is active", bool(_a) and not _a["indefinite"])
check("a timed pause is not indefinite", _a.get("until") is not None)
check("status line shows a countdown", "paused" in W.pause_status_line(_cfgp))
_w = W.end_pause(_cfgp)
check("resume closes the window", _w and _w["ended_by"] == "resume")
check("resume clears the active pause", W.active_pause(_cfgp) is None)
check("the window is appended to history",
      (_ps / "pauses.jsonl").read_text().count("\n") == 1)

# an indefinite pause must still expire at the end of the workday
_r2 = W.start_pause(_cfgp)
check("an indefinite pause is flagged as such", _r2["indefinite"])
_now2 = datetime.now()
_ws, _we = W.window_bounds(_cfgp, _now2.date())
_inside = _now2.isoweekday() in _cfgp["work_days"] and _ws <= _now2 < _we
_expect = _we if _inside else W.next_window_start(_cfgp, _now2)
check("an indefinite pause ends at today's close inside hours, else at the next start",
      _r2["until"] is not None
      and datetime.fromisoformat(_r2["until"]) == _expect,
      f"{_r2['until']} vs {_expect}")
W.end_pause(_cfgp)

# an indefinite pause is ALWAYS bounded, whatever time of day it is started
check("next_window_start skips to tomorrow after hours",
      W.next_window_start(_cfgp, datetime(2026, 8, 13, 17, 53)) == datetime(2026, 8, 14, 8, 0),
      str(W.next_window_start(_cfgp, datetime(2026, 8, 13, 17, 53))))
check("next_window_start skips the weekend",
      W.next_window_start(_cfgp, datetime(2026, 8, 15, 12, 0)) == datetime(2026, 8, 17, 8, 0),
      str(W.next_window_start(_cfgp, datetime(2026, 8, 15, 12, 0))))
_r3 = W.start_pause(_cfgp)
_u3 = _r3["until"]
check("an indefinite pause started after hours is still bounded", _u3 is not None, str(_u3))
if _u3:
    # The real invariant is not a fixed number of hours — a Friday evening pause
    # legitimately runs to Monday, because the sampler is idle all weekend anyway.
    # It is that the bound never reaches past the moment tracking should resume:
    # bounding to the next window's END instead would swallow a whole workday.
    _u3d = datetime.fromisoformat(_u3)
    _now3 = datetime.fromisoformat(_r3["started"])
    _ws3, _we3 = W.window_bounds(_cfgp, _now3.date())
    _inside3 = _now3.isoweekday() in _cfgp["work_days"] and _ws3 <= _now3 < _we3
    check("an indefinite pause never runs past the next work window start",
          _u3d == (_we3 if _inside3 else W.next_window_start(_cfgp, _now3)),
          f"{_u3d} (started {_now3}, inside hours: {_inside3})")
W.end_pause(_cfgp)

# "end of day" is the calendar day, and a separate mode waits for the next window
class _A:  # a stand-in for argparse's namespace
    def __init__(self, **kw):
        self.minutes = None; self.until = None; self.rest_of_day = False
        self.next_period = False; self.reason = ""; self.quiet = True
        self.__dict__.update(kw)


(_ps / "pause.json").unlink(missing_ok=True)
W.cmd_pause(_cfgp, _A(rest_of_day=True))
_mid = W.active_pause(_cfgp)
check("pause until end of day means midnight tonight",
      _mid and datetime.fromisoformat(_mid["until"])
      == datetime.combine(D.today() + timedelta(days=1), datetime.min.time()),
      str((_mid or {}).get("until")))
check("midnight is reachable at any hour, so it never has to refuse", _mid is not None)
W.end_pause(_cfgp)

W.cmd_pause(_cfgp, _A(next_period=True))
_np = W.active_pause(_cfgp)
check("pause until the next period ends when that window opens",
      _np and datetime.fromisoformat(_np["until"])
      == W.next_window_start(_cfgp, datetime.fromisoformat(_np["started"])),
      str((_np or {}).get("until")))
W.end_pause(_cfgp)
check("the menu bar label can be switched off",
      "show_label" in Path(W.__file__).read_text())

# an expired pause retires itself on sight rather than lingering
(_ps / "pause.json").write_text(json.dumps(
    {"started": "2026-08-12T09:00:00", "until": "2026-08-12T09:30:00",
     "indefinite": False, "reason": "personal"}))
check("an expired pause is not reported active", W.active_pause(_cfgp, _N) is None)
check("an expired pause is retired into history",
      "09:30:00" in (_ps / "pauses.jsonl").read_text())
check("an expired pause is closed at its until, not at discovery",
      json.loads((_ps / "pauses.jsonl").read_text().strip().splitlines()[-1])["ended_by"]
      == "expiry")

# containment, not overlap
_win(datetime(2026, 8, 12, 12, 0), datetime(2026, 8, 12, 13, 0))
_pw = W.pause_windows(_cfgp, D(2026, 8, 12))
check("the window is found for that day", len(_pw) == 1, str(_pw))
check("an item inside the window is paused",
      W.in_pause(_pw, datetime(2026, 8, 12, 12, 30)))
check("an item before the window is not",
      not W.in_pause(_pw, datetime(2026, 8, 12, 11, 59)))
check("an item after the window is not",
      not W.in_pause(_pw, datetime(2026, 8, 12, 13, 1)))
check("a span straddling the window survives — containment, not overlap",
      not W.in_pause(_pw, datetime(2026, 8, 12, 11, 0), datetime(2026, 8, 12, 14, 0)))
check("a span wholly inside is dropped",
      W.in_pause(_pw, datetime(2026, 8, 12, 12, 10), datetime(2026, 8, 12, 12, 50)))

_items = [{"at": "2026-08-12T11:30:00", "id": "before"},
          {"at": "2026-08-12T12:30:00", "id": "inside"},
          {"at": "2026-08-12T13:30:00", "id": "after"}]
_kept, _dropped = W.drop_paused(_pw, _items, "at")
check("drop_paused removes only the contained item",
      [i["id"] for i in _kept] == ["before", "after"] and _dropped == 1)
_sessions = [{"start": "2026-08-12T11:00:00", "end": "2026-08-12T14:00:00", "id": "long"},
             {"start": "2026-08-12T12:05:00", "end": "2026-08-12T12:50:00", "id": "short"}]
_kept2, _d2 = W.drop_paused(_pw, _sessions, "start", "end")
check("a long session straddling a pause is kept",
      [i["id"] for i in _kept2] == ["long"] and _d2 == 1)
check("drop_paused is a no-op with no windows",
      W.drop_paused([], _items, "at") == (_items, 0))
check("timezone-aware timestamps are handled",
      W._naive("2026-08-12T12:30:00+00:00") is not None
      and W._naive("2026-08-12T12:30:00") == datetime(2026, 8, 12, 12, 30))
check("an unparseable timestamp is kept rather than dropped",
      W.drop_paused(_pw, [{"at": "not a date"}], "at")[1] == 0)

# paused minutes are discounted so health does not cry wolf
_lo, _hi = datetime(2026, 8, 12, 8, 0), datetime(2026, 8, 12, 17, 0)
check("paused minutes are measured within the window",
      abs(W.paused_minutes(_cfgp, D(2026, 8, 12), _lo, _hi) - 60.0) < 0.1,
      str(W.paused_minutes(_cfgp, D(2026, 8, 12), _lo, _hi)))
check("paused minutes clip to the window asked for",
      abs(W.paused_minutes(_cfgp, D(2026, 8, 12), datetime(2026, 8, 12, 12, 30),
                           _hi) - 30.0) < 0.1)
(_ps / "raw" / "2026-08-12").mkdir(parents=True, exist_ok=True)
(_ps / "raw" / "2026-08-12" / "activity.jsonl").write_text(
    "".join('{"ts":"x"}\n' for _ in range(480)))
_h = W.sampler_health(_cfgp, D(2026, 8, 12), now=datetime(2026, 8, 12, 17, 0))
check("health subtracts the paused hour from expected",
      _h and _h["paused_minutes"] == 60.0 and _h["expected"] == 480, str(_h))
check("a full day minus a paused hour reads healthy", _h and _h["ok"], str(_h))

# the recoverability guarantee: trimming a window puts the work back
_cfgq = json.loads(json.dumps(_cfgp))
_gitish = [{"at": "2026-08-12T12:30:00", "sha": "aaa"},
           {"at": "2026-08-12T15:00:00", "sha": "bbb"}]
check("with the pause in place the contained commit is hidden",
      len(W.drop_paused(W.pause_windows(_cfgq, D(2026, 8, 12)), _gitish, "at")[0]) == 1)
(_ps / "pauses.jsonl").write_text("")          # the user trims the bad window
check("trimming the window restores the hidden commit — loss is reversible",
      len(W.drop_paused(W.pause_windows(_cfgq, D(2026, 8, 12)), _gitish, "at")[0]) == 2)

check("every timestamped collector consults the pause history",
      Path(W.__file__).read_text().count("drop_paused(") >= 7)
check("mis_board is knowingly unfiltered (rows carry a date, not a time)",
      "date_column" in (CFG.get("mis_board") or {}))

# ----------------------------------------------------------------- tokenwise --
_tw_day = D(2026, 9, 4)
check("tokenwise window is expressed in UTC",
      W._utc_str(datetime(2026, 9, 4, 8, 0)) ==
      datetime(2026, 9, 4, 8, 0).astimezone(W.timezone.utc).strftime("%Y-%m-%dT%H:%M:%S"))
check("fmt_tokens rounds the way the ledger does",
      (W.fmt_tokens(950), W.fmt_tokens(1234), W.fmt_tokens(950_000),
       W.fmt_tokens(2_430_000_000)) == ("950", "1K", "950K", "2.43B"))
check("tokenwise disabled reports so",
      W.collect_tokenwise({**CFG, "tokenwise": {**CFG.get("tokenwise", {}), "enabled": False}},
                          _tw_day)["reason"] == "disabled")
_missing = W.collect_tokenwise({**CFG, "tokenwise": {"enabled": True, "dir": str(TMP)}}, _tw_day)
check("tokenwise missing ledger is unavailable, not an error",
      _missing.get("available") is False and "not found" in _missing.get("reason", ""))
check("the summary prompt asks for a usage block and forbids tracker rows from it",
      "`Claude Code usage`" in W.SUMMARY_PROMPT and "never turn them" in W.SUMMARY_PROMPT)
_tw = W.collect_tokenwise(CFG, _tw_day, ingest=False)
if not (_tw.get("available") and _tw.get("turns")):
    skip.append(f"tokenwise: no ledger turns on {_tw_day} ({_tw.get('reason', 'empty')})")
else:
    check("tokenwise count mirrors main-thread turns", _tw["count"] == _tw["turns"])
    check("tokenwise sessions are ranked by cache-read",
          [s["cache_read"] for s in _tw["sessions"]] ==
          sorted((s["cache_read"] for s in _tw["sessions"]), reverse=True))
    _thr = W.tokenwise_cfg(CFG)["long_session_turns"]
    check("tokenwise long flag follows the configured threshold",
          all(s["long"] == (s["turns_total"] >= _thr) for s in _tw["sessions"]))
    check("tokenwise a session never has more turns today than in total",
          all(s["turns_today"] <= s["turns_total"] for s in _tw["sessions"]))
    check("tokenwise session_file joins to the claude_code collector's key",
          all(s["session_file"].endswith(".jsonl") for s in _tw["sessions"]))
    check("tokenwise projects carry no home prefix",
          all(not s["project"].startswith("Users-") for s in _tw["sessions"]),
          str([s["project"] for s in _tw["sessions"]])[:120])
    check("tokenwise share over 400K is a percentage",
          0 <= _tw["context"]["share_over_400k"] <= 100)
    check("tokenwise cache-read total equals the model breakdown",
          _tw["tokens"]["cache_read"] == sum(m["cache_read"] for m in _tw["by_model"].values()))
    check("tokenwise advice appears only with a long session",
          bool(_tw.get("advice")) == any(s["long"] for s in _tw["sessions"]))
    # rendering: the fallback and the notification headline pick the block up
    _raw = Path.home() / ".worklog" / "raw" / _tw_day.isoformat() / "digest.json"
    if _raw.is_file():
        _dg = json.loads(_raw.read_text())
        _dg["sources"]["tokenwise"] = _tw
        _md = W.render_fallback(CFG, _dg)
        check("fallback renders the usage block", "**Claude Code usage**" in _md)
        check("fallback names long sessions with a /clear hint",
              ("/clear" in _md) == any(s["long"] for s in _tw["sessions"]))
        check("headline mentions cache-read", "cache-read" in W.digest_headline(_dg))
    else:
        skip.append(f"tokenwise: no stored digest for {_tw_day} to render")
    _wk = W.tokenwise_week(CFG, [_tw_day])
    check("weekly usage lines start with the heading",
          bool(_wk) and _wk[0] == "**Claude Code usage — week**" and len(_wk) >= 2, str(_wk)[:120])

print(f"\n{len(ok)} passed, {len(fail)} failed, {len(skip)} skipped\n")
for s in skip:
    print(f"  SKIP  {s}")
for f in fail:
    print(f"  FAIL  {f}")
if "-v" in sys.argv:
    for o in ok:
        print(f"  PASS  {o}")
sys.exit(1 if fail else 0)
