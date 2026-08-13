#!/usr/bin/env python3
"""
worklog — passive workday tracking + end-of-day markdown summary (macOS).

Subcommands:
  sample                 Take one activity sample. Exits silently outside work hours.
  report [--date D]      Build the digest and write the markdown work log.
  repair [--date D]      Redo days whose log is missing or was written without the LLM.
  pause / resume         Stop and start recording during personal time.
  status                 Current state as JSON, for the menu bar.
  weekly [--date D]      Write the week's merged MIS table (YYYY-Www.md).
  digest [--date D]      Print the raw JSON digest (no summarization). For debugging.
  doctor                 Check paths, permissions and data sources.

Config: ~/.worklog/config.json
"""

from __future__ import annotations

import argparse
import fcntl
import json
import os
import re
import shutil
import subprocess
import sys
import time
from collections import defaultdict
from datetime import datetime, date as date_cls, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

SELF_DIR = Path(__file__).resolve().parent
DEFAULT_STATE = Path("~/.worklog").expanduser()


# --------------------------------------------------------------------------- #
# config / util
# --------------------------------------------------------------------------- #

def expand(p: str) -> Path:
    return Path(os.path.expanduser(str(p)))


def load_config() -> dict:
    for cand in (DEFAULT_STATE / "config.json", SELF_DIR.parent / "config.json"):
        if cand.is_file():
            with open(cand) as fh:
                return json.load(fh)
    die(f"No config.json found (looked in {DEFAULT_STATE} and {SELF_DIR.parent})")


def die(msg: str, code: int = 1):
    print(f"worklog: {msg}", file=sys.stderr)
    sys.exit(code)


def log_line(cfg: dict, msg: str):
    logs = expand(cfg["state_dir"]) / "logs"
    logs.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().isoformat(timespec="seconds")
    with open(logs / "worklog.log", "a") as fh:
        fh.write(f"{stamp} {msg}\n")


def run(cmd, timeout=30, cwd=None) -> tuple[int, str, str]:
    """Run a command, never raise. Returns (rc, stdout, stderr)."""
    try:
        p = subprocess.run(
            cmd, capture_output=True, text=True, timeout=timeout, cwd=cwd,
            errors="replace",
        )
        return p.returncode, p.stdout.strip(), p.stderr.strip()
    except subprocess.TimeoutExpired:
        return 124, "", "timeout"
    except (FileNotFoundError, OSError) as exc:
        return 127, "", str(exc)


def osascript(script: str, timeout=30) -> tuple[int, str, str]:
    return run(["osascript", "-e", script], timeout=timeout)


def parse_hhmm(s: str) -> tuple[int, int]:
    h, m = s.split(":")
    return int(h), int(m)


def window_bounds(cfg: dict, day: date_cls) -> tuple[datetime, datetime]:
    sh, sm = parse_hhmm(cfg["work_hours"]["start"])
    eh, em = parse_hhmm(cfg["work_hours"]["end"])
    return (
        datetime.combine(day, datetime.min.time()).replace(hour=sh, minute=sm),
        datetime.combine(day, datetime.min.time()).replace(hour=eh, minute=em),
    )


def prev_workday(cfg: dict, day: date_cls) -> date_cls | None:
    for i in range(1, 8):
        p = day - timedelta(days=i)
        if p.isoweekday() in cfg["work_days"]:
            return p
    return None


def is_work_time(cfg: dict, now: datetime) -> bool:
    if now.isoweekday() not in cfg["work_days"]:
        return False
    start, end = window_bounds(cfg, now.date())
    return start <= now <= end


def fmt_dur(minutes: float) -> str:
    minutes = int(round(minutes))
    if minutes < 60:
        return f"{minutes}m"
    h, m = divmod(minutes, 60)
    return f"{h}h" if m == 0 else f"{h}h {m:02d}m"


# --------------------------------------------------------------------------- #
# pause: stop recording during personal time
#
# A pause hides work, so the windows are kept as durable, editable data rather
# than a transient flag. Collectors filter BY READING this history and a repair
# re-collects from source, so trimming a line here and re-running the repair puts
# any wrongly-dropped work back. That is what makes forgetting to resume
# recoverable instead of permanent.
# --------------------------------------------------------------------------- #

def pause_paths(cfg: dict) -> tuple[Path, Path]:
    state = expand(cfg["state_dir"])
    return state / "pause.json", state / "pauses.jsonl"


def active_pause(cfg: dict, now: datetime | None = None) -> dict | None:
    """The pause in force right now, or None. Expired ones are retired on sight."""
    now = now or datetime.now()
    cur, _ = pause_paths(cfg)
    if not cur.is_file():
        return None
    try:
        p = json.loads(cur.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError):
        return None
    until = p.get("until")
    if until:
        try:
            if now >= datetime.fromisoformat(until):
                end_pause(cfg, now, expired=True)
                return None
        except ValueError:
            return None
    return p


def pause_windows(cfg: dict, day: date_cls) -> list[tuple[datetime, datetime]]:
    """Every paused interval touching `day`, including one still running."""
    _, hist = pause_paths(cfg)
    out = []
    if hist.is_file():
        for line in hist.read_text(errors="replace").splitlines():
            if not line.strip():
                continue
            try:
                w = json.loads(line)
                s = datetime.fromisoformat(w["started"])
                e = datetime.fromisoformat(w["ended"])
            except (json.JSONDecodeError, KeyError, ValueError):
                continue
            if s.date() <= day <= e.date():
                out.append((s, e))
    cur = active_pause(cfg)
    if cur:
        try:
            s = datetime.fromisoformat(cur["started"])
            e = (datetime.fromisoformat(cur["until"]) if cur.get("until")
                 else datetime.now())
            if s.date() <= day <= e.date():
                out.append((s, e))
        except (KeyError, ValueError):
            pass
    return sorted(out)


def in_pause(windows: list, start, end=None) -> bool:
    """True when [start, end] sits ENTIRELY inside one paused window.

    Containment, not overlap: a three-hour session must not disappear because of a
    fifteen-minute pause in the middle of it.
    """
    if not windows or start is None:
        return False
    a = start if isinstance(start, datetime) else None
    if a is None:
        return False
    b = end if isinstance(end, datetime) else a
    return any(ws <= a and b <= we for ws, we in windows)


def _naive(ts) -> datetime | None:
    """Parse a timestamp to naive local time, whatever shape it arrives in."""
    if isinstance(ts, datetime):
        return ts.astimezone().replace(tzinfo=None) if ts.tzinfo else ts
    if not ts:
        return None
    try:
        d = datetime.fromisoformat(str(ts).replace("Z", "+00:00"))
    except ValueError:
        return None
    return d.astimezone().replace(tzinfo=None) if d.tzinfo else d


def drop_paused(windows: list, items: list, start_key: str,
                end_key: str | None = None) -> tuple[list, int]:
    """Remove items lying entirely within a paused window. Returns (kept, dropped)."""
    if not windows or not items:
        return items, 0
    kept = []
    for it in items:
        s = _naive(it.get(start_key))
        e = _naive(it.get(end_key)) if end_key else s
        if s is not None and in_pause(windows, s, e or s):
            continue
        kept.append(it)
    return kept, len(items) - len(kept)


def paused_minutes(cfg: dict, day: date_cls, lo: datetime, hi: datetime) -> float:
    """Paused minutes inside [lo, hi] — used so health checks stay honest."""
    total = 0.0
    for ws, we in pause_windows(cfg, day):
        s, e = max(ws, lo), min(we, hi)
        if e > s:
            total += (e - s).total_seconds() / 60.0
    return total


def start_pause(cfg: dict, minutes: float | None = None,
                until: datetime | None = None, reason: str = "") -> dict:
    """Begin a pause. No `minutes`/`until` means indefinite, bounded by the workday."""
    now = datetime.now()
    if minutes:
        until = now + timedelta(minutes=float(minutes))
    if until is None:
        # An open-ended pause still expires at the end of the workday: the worst case
        # then costs one day, not a silently untracked week.
        _, day_end = window_bounds(cfg, now.date())
        until = day_end if day_end > now else None
        indefinite = True
    else:
        indefinite = False
    rec = {"started": now.isoformat(timespec="seconds"),
           "until": until.isoformat(timespec="seconds") if until else None,
           "indefinite": indefinite, "reason": reason or "personal",
           "last_reminder": now.isoformat(timespec="seconds")}
    cur, _ = pause_paths(cfg)
    cur.parent.mkdir(parents=True, exist_ok=True)
    cur.write_text(json.dumps(rec), encoding="utf-8")
    log_line(cfg, f"pause: started {rec['started']} until {rec['until'] or 'workday end'}"
                  f"{' (indefinite)' if indefinite else ''}")
    return rec


def remind_if_indefinite(cfg: dict, paused: dict, now: datetime) -> None:
    """Nag while an open-ended pause runs, so it cannot be forgotten quietly."""
    if not (paused.get("indefinite") or not paused.get("until")):
        return
    every = float((cfg.get("pause") or {}).get("reminder_minutes", 30))
    if every <= 0:
        return
    try:
        last = datetime.fromisoformat(paused.get("last_reminder") or paused["started"])
    except (KeyError, ValueError, TypeError):
        return
    if (now - last).total_seconds() / 60.0 < every:
        return
    try:
        since = fmt_dur((now - datetime.fromisoformat(paused["started"]))
                        .total_seconds() / 60.0)
    except (KeyError, ValueError):
        since = "a while"
    notify(cfg, "Work log still paused", f"Nothing recorded for {since}.",
           subtitle="run `worklog resume` to start tracking again")
    paused["last_reminder"] = now.isoformat(timespec="seconds")
    cur, _ = pause_paths(cfg)
    try:
        cur.write_text(json.dumps(paused), encoding="utf-8")
    except OSError:
        pass


def end_pause(cfg: dict, now: datetime | None = None, expired: bool = False) -> dict | None:
    """Close the active pause and append the window to the durable history."""
    now = now or datetime.now()
    cur, hist = pause_paths(cfg)
    if not cur.is_file():
        return None
    try:
        p = json.loads(cur.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError):
        cur.unlink(missing_ok=True)
        return None
    started = p.get("started")
    ended = now
    if expired and p.get("until"):
        try:
            ended = datetime.fromisoformat(p["until"])
        except ValueError:
            pass
    window = {"started": started,
              "ended": ended.isoformat(timespec="seconds"),
              "reason": p.get("reason", "personal"),
              "ended_by": "expiry" if expired else "resume"}
    hist.parent.mkdir(parents=True, exist_ok=True)
    with open(hist, "a", encoding="utf-8") as fh:
        fh.write(json.dumps(window) + "\n")
    cur.unlink(missing_ok=True)
    log_line(cfg, f"pause: ended {window['ended']} by {window['ended_by']}")
    return window


# --------------------------------------------------------------------------- #
# redaction
# --------------------------------------------------------------------------- #

_SECRET_RULES = [
    (re.compile(r"(?i)\b(api[_-]?key|apikey|token|secret|password|passwd|pwd|bearer|"
                r"auth[_-]?token|access[_-]?token|client[_-]?secret)\b(\s*[=:]\s*|\s+)"
                r"(['\"]?)([^\s'\"]{6,})\3"), r"\1\2«redacted»"),
    (re.compile(r"\bAKIA[0-9A-Z]{16}\b"), "«redacted-aws-key»"),
    (re.compile(r"\bsk-ant-[A-Za-z0-9_\-]{16,}"), "«redacted-anthropic-key»"),
    (re.compile(r"\bsk-[A-Za-z0-9]{20,}"), "«redacted-key»"),
    (re.compile(r"\bgh[pousr]_[A-Za-z0-9]{16,}"), "«redacted-github-token»"),
    (re.compile(r"\bxox[baprs]-[A-Za-z0-9\-]{10,}"), "«redacted-slack-token»"),
    (re.compile(r"\bey[A-Za-z0-9_\-]{10,}\.[A-Za-z0-9_\-]{10,}\.[A-Za-z0-9_\-]{10,}"),
     "«redacted-jwt»"),
    (re.compile(r"(?i)(https?://)([^/\s:@]+):([^/\s@]+)@"), r"\1\2:«redacted»@"),
]


class Redactor:
    def __init__(self, cfg: dict):
        rc = cfg.get("redaction", {})
        self.enabled = rc.get("enabled", True)
        self.rules = list(_SECRET_RULES)
        for pat in rc.get("extra_patterns", []):
            try:
                self.rules.append((re.compile(pat), "«redacted»"))
            except re.error:
                pass

    def __call__(self, text: str) -> str:
        if not self.enabled or not text:
            return text
        for rx, repl in self.rules:
            text = rx.sub(repl, text)
        return text


def notify(cfg: dict, title: str, message: str, subtitle: str = "",
           open_path: str | None = None) -> None:
    """Post a macOS notification. Never raises; failure is logged and ignored."""
    ncfg = cfg.get("notifications", {})
    if not ncfg.get("enabled", True):
        return

    tn = shutil.which("terminal-notifier")
    if tn:
        cmd = [tn, "-title", title, "-message", message, "-group", "com.workbuddy"]
        icon = ncfg.get("icon") or "~/.worklog/assets/icon-256.png"
        paused_icon = ncfg.get("paused_icon") or "~/.worklog/assets/icon-paused-256.png"
        pick = paused_icon if "paus" in f"{title} {message}".lower() else icon
        ip = expand(pick)
        if ip.is_file():
            # -contentImage rides alongside the notification; -appIcon replaces the
            # terminal-notifier icon, which is what makes it look like our app
            cmd += ["-contentImage", str(ip), "-appIcon", str(ip)]
        if subtitle:
            cmd += ["-subtitle", subtitle]
        if ncfg.get("sound"):
            cmd += ["-sound", ncfg["sound"]]
        if open_path and ncfg.get("click_opens_log", True):
            cmd += ["-execute", f'open "{open_path}"']
        rc, _, err = run(cmd, timeout=20)
        if rc == 0:
            return
        log_line(cfg, f"notify: terminal-notifier rc={rc} {err[:160]}")

    def esc(s: str) -> str:
        return s.replace("\\", "\\\\").replace('"', '\\"')

    script = f'display notification "{esc(message)}" with title "{esc(title)}"'
    if subtitle:
        script += f' subtitle "{esc(subtitle)}"'
    if ncfg.get("sound"):
        script += f' sound name "{esc(ncfg["sound"])}"'
    rc, _, err = osascript(script, timeout=20)
    if rc != 0:
        log_line(cfg, f"notify: osascript rc={rc} {err[:160]}")


def digest_headline(d: dict) -> str:
    """Short 'what got captured' line for the notification body."""
    s = d["sources"]
    bits = []
    if s["git"].get("total_commits"):
        n = s["git"]["total_commits"]
        bits.append(f"{n} commit{'s' if n != 1 else ''}")
    n = max(len(s["calendar"].get("events") or []),
            s.get("unjoined_meetings", {}).get("count") or 0,
            s.get("meetings_attended", {}).get("count") or 0)
    if n:
        bits.append(f"{n} meeting{'s' if n != 1 else ''}")
    n = s.get("mis_board", {}).get("count") or 0
    if n:
        bits.append(f"{n} board task{'s' if n != 1 else ''}")
    if s["claude_code"].get("session_count"):
        n = s["claude_code"]["session_count"]
        bits.append(f"{n} CC session{'s' if n != 1 else ''}")
    if s["cloud"].get("count"):
        bits.append(f"{s['cloud']['count']} file{'s' if s['cloud']['count'] != 1 else ''}")
    if s["activity"].get("active_duration"):
        bits.append(f"{s['activity']['active_duration']} active")
    return " · ".join(bits) if bits else "no activity captured"


# --------------------------------------------------------------------------- #
# sampler
# --------------------------------------------------------------------------- #

FRONT_APP_SCRIPT = '''
tell application "System Events"
  set appName to ""
  set winTitle to ""
  set bundleId to ""
  try
    set appName to name of first application process whose frontmost is true
  end try
  try
    set bundleId to bundle identifier of (first application process whose frontmost is true)
  end try
  try
    set winTitle to name of front window of (first application process whose frontmost is true)
  end try
end tell
return appName & "\\t" & winTitle & "\\t" & bundleId
'''


EXCLUDED_APP = "(excluded)"

_APP_DIRS = ("/System/Applications", "/System/Applications/Utilities",
             "/Applications", "/Applications/Utilities",
             "/System/Library/CoreServices")
_BUNDLE_CACHE: dict = {}


def bundle_for_app(name: str) -> str | None:
    """Bundle id for an app name, read off disk.

    Needed because samples written before the sampler recorded the bundle id only
    have a name, and because a vendor rule has to apply to those too.
    """
    key = (name or "").strip()
    if not key or key.startswith("("):
        return None
    if key in _BUNDLE_CACHE:
        return _BUNDLE_CACHE[key]
    bid = None
    for d in _APP_DIRS:
        p = Path(d) / f"{key}.app" / "Contents" / "Info.plist"
        if p.is_file():
            try:
                import plistlib
                bid = plistlib.loads(p.read_bytes()).get("CFBundleIdentifier")
            except Exception:
                bid = None
            break
    _BUNDLE_CACHE[key] = bid
    return bid


def is_excluded_app(acfg: dict, app: str, bundle: str | None = None) -> bool:
    """True when this app must leave no trace in the log.

    Two rules. An explicit name in `exclude_apps`, matched case-insensitively so
    `zoom.us` and `Zoom Workplace` cannot slip through. And a vendor rule: a bundle
    id under any `vendor_exclude_prefixes` is excluded unless it is listed in
    `vendor_allow_bundles` — that way every Apple app is covered, including ones
    never opened before, without maintaining a list of app names.
    """
    name = (app or "").strip()
    if any(name.lower() == str(x).strip().lower() for x in acfg.get("exclude_apps") or []):
        return True
    prefixes = acfg.get("vendor_exclude_prefixes") or []
    if not prefixes:
        return False
    bid = (bundle or "").strip() or bundle_for_app(name)
    if not bid:
        return False
    if not any(bid.lower().startswith(str(p).lower()) for p in prefixes):
        return False
    allow = {str(a).strip().lower() for a in acfg.get("vendor_allow_bundles") or []}
    return bid.lower() not in allow


def idle_seconds() -> int:
    rc, out, _ = run(["ioreg", "-c", "IOHIDSystem"], timeout=10)
    if rc != 0:
        return 0
    m = re.search(r'"HIDIdleTime"\s*=\s*(\d+)', out)
    return int(int(m.group(1)) / 1_000_000_000) if m else 0


MEETING_WINDOWS_SCRIPT = '''
tell application "System Events"
  set out to ""
  repeat with pn in {apps}
    set p to pn as string
    if exists process p then
      tell process p
        repeat with w in windows
          try
            set out to out & p & tab & (name of w) & linefeed
          end try
        end repeat
      end tell
    end if
  end repeat
end tell
return out
'''


def meeting_windows_now(cfg: dict) -> list[list[str]]:
    """Every open window of the meeting apps, as [app, title] pairs.

    Frontmost sampling alone undercounts badly: a call you tab away from stops
    accruing time. A meeting window exists for as long as you are in the call,
    whether or not it has focus.
    """
    apps = cfg["activity"].get("meeting_apps") or []
    if not apps:
        return []
    literal = "{" + ", ".join('"' + a.replace('"', '\\"') + '"' for a in apps) + "}"
    rc, out, err = osascript(MEETING_WINDOWS_SCRIPT.format(apps=literal), timeout=20)
    if rc != 0:
        log_line(cfg, f"sample: meeting-window probe failed rc={rc} {err[:160]}")
        return []
    pairs = []
    for line in out.splitlines():
        app, _, title = line.partition("\t")
        if app and title.strip():
            pairs.append([app, title.strip()])
    return pairs


SCHEDULED_MENU_SCRIPT = '''
tell application "System Events"
  if not (exists process "{app}") then return ""
  tell process "{app}"
    repeat with mb in menu bars
      repeat with mbi in (menu bar items of mb)
        try
          if (description of mbi) is "status menu" then return (title of mbi)
        end try
      end repeat
    end repeat
  end tell
end tell
return ""
'''

# Outlook's status menu is the only readable calendar signal on this machine, and
# it is a weak one. Measured 2026-07-29: it names the next meeting you have NOT
# joined. Joining a meeting makes it skip straight to the next one, and the user
# reports some booked meetings never appear at all. So `Now:` means "booked, running,
# and not joined here" — NOT "what was scheduled". Everything else was ruled out:
# New Outlook has no AppleScript event store, no Exchange account is registered
# (MDM blocks adding one) so Calendar.app and icalBuddy are empty, HxStore.hxd holds
# subjects but no decodable timestamps near them, and Outlook's meeting reminders
# never reach the Notification Center database.
_MEETING_SUFFIXES = (" Microsoft Teams Meeting", " Teams Meeting", " Zoom Meeting",
                     " Webex Meeting", " Skype Meeting", " Google Meet")


def parse_scheduled(title: str) -> str | None:
    """Subject of the meeting Outlook says is on NOW, or None.

    Only the `Now:` prefix counts — `Next:` announces something that has not
    started and must never be recorded as time spent.
    """
    t = " ".join((title or "").split())
    if not t.lower().startswith("now:"):
        return None
    t = t[4:].strip()
    for suffix in _MEETING_SUFFIXES:
        if t.lower().endswith(suffix.lower()):
            t = t[: -len(suffix)].strip()
            break
    return t or None


def scheduled_meeting_now(cfg: dict) -> str | None:
    scfg = cfg.get("unjoined_meetings") or {}
    if not scfg.get("enabled", True):
        return None
    app = scfg.get("app", "Microsoft Outlook")
    rc, out, err = osascript(SCHEDULED_MENU_SCRIPT.format(app=app), timeout=20)
    if rc != 0:
        log_line(cfg, f"sample: scheduled-menu probe failed rc={rc} {err[:160]}")
        return None
    return parse_scheduled(out)


def cmd_sample(cfg: dict, args) -> int:
    now = datetime.now()
    if not args.force and not is_work_time(cfg, now):
        return 0
    if not cfg["activity"]["enabled"]:
        return 0
    # Paused: write nothing at all. No row means no app time, no window title and no
    # meeting presence for this minute; the pause history is the only record of it.
    paused = active_pause(cfg, now)
    if paused:
        remind_if_indefinite(cfg, paused, now)
        return 0

    idle = idle_seconds()
    bundle = ""
    if idle >= cfg["idle_threshold_seconds"]:
        app, title = "(idle)", ""
    else:
        rc, out, err = osascript(FRONT_APP_SCRIPT, timeout=15)
        if rc != 0:
            log_line(cfg, f"sample: osascript failed rc={rc} {err[:200]}")
            app, title = "(unknown)", ""
        else:
            parts = out.split("\t")
            app = parts[0].strip() or "(unknown)"
            title = parts[1].strip() if len(parts) > 1 else ""
            bundle = parts[2].strip() if len(parts) > 2 else ""

    acfg = cfg["activity"]
    if is_excluded_app(acfg, app, bundle):
        # the name is dropped too, not just the title: an excluded app must leave no
        # trace, and "Spotify (excluded)" in a work log is still a record of Spotify
        title = ""
        app = EXCLUDED_APP
    for pat in acfg["exclude_title_patterns"]:
        try:
            if title and re.search(pat, title):
                title = "(excluded)"
                break
        except re.error:
            pass

    redact = Redactor(cfg)
    rec = {
        "ts": now.isoformat(timespec="seconds"),
        "app": app,
        "title": redact(title)[:300],
        "idle": idle,
    }
    if bundle and app != EXCLUDED_APP:
        rec["bundle"] = bundle[:120]

    # recorded even while idle: sitting in a call listening is exactly the case
    # frontmost sampling loses
    windows = []
    for wapp, wtitle in meeting_windows_now(cfg):
        if is_excluded_app(acfg, wapp):
            continue
        for pat in acfg["exclude_title_patterns"]:
            try:
                if re.search(pat, wtitle):
                    wtitle = "(excluded)"
                    break
            except re.error:
                pass
        windows.append([wapp, redact(wtitle)[:300]])
    rec["meeting_windows"] = windows

    sched = scheduled_meeting_now(cfg)
    if sched:
        for pat in acfg["exclude_title_patterns"]:
            try:
                if re.search(pat, sched):
                    sched = None
                    break
            except re.error:
                pass
    rec["scheduled_now"] = redact(sched)[:300] if sched else None
    raw_dir = expand(cfg["state_dir"]) / "raw" / now.strftime("%Y-%m-%d")
    raw_dir.mkdir(parents=True, exist_ok=True)
    with open(raw_dir / "activity.jsonl", "a") as fh:
        fh.write(json.dumps(rec, ensure_ascii=False) + "\n")
    return 0


# --------------------------------------------------------------------------- #
# collector: activity rollup
# --------------------------------------------------------------------------- #

def collect_activity(cfg: dict, day: date_cls) -> dict:
    path = expand(cfg["state_dir"]) / "raw" / day.strftime("%Y-%m-%d") / "activity.jsonl"
    if not path.is_file():
        return {"available": False, "reason": "no samples recorded for this day"}

    samples = []
    with open(path, errors="replace") as fh:
        for line in fh:
            line = line.strip()
            if not line:
                continue
            try:
                rec = json.loads(line)
                rec["_dt"] = datetime.fromisoformat(rec["ts"])
                samples.append(rec)
            except (json.JSONDecodeError, KeyError, ValueError):
                continue
    if not samples:
        return {"available": False, "reason": "sample file empty or unreadable"}

    samples.sort(key=lambda r: r["_dt"])
    per_app = defaultdict(float)
    titles = defaultdict(lambda: defaultdict(float))
    idle_min = 0.0
    # A sample only ever credits up to this much time. Anything longer means the
    # machine was asleep or the agent was not running — not time spent in that app.
    cap = float(cfg["activity"].get("sample_gap_cap_minutes", 2))

    for i, rec in enumerate(samples):
        if i + 1 < len(samples):
            gap = (samples[i + 1]["_dt"] - rec["_dt"]).total_seconds() / 60.0
        else:
            gap = 1.0
        gap = min(max(gap, 0.0), cap)
        if rec["app"] == EXCLUDED_APP or is_excluded_app(
                cfg["activity"], rec["app"], rec.get("bundle")):
            # neither active nor idle: excluded time is simply not part of the log,
            # so the day's totals are lower rather than silently reattributed
            continue
        if rec["app"] == "(idle)":
            idle_min += gap
            continue
        per_app[rec["app"]] += gap
        if rec.get("title"):
            titles[rec["app"]][rec["title"]] += gap

    acfg = cfg["activity"]
    apps = []
    for app, mins in sorted(per_app.items(), key=lambda kv: -kv[1]):
        if mins < acfg["min_minutes_to_report"]:
            continue
        top = sorted(titles[app].items(), key=lambda kv: -kv[1])[: acfg["max_titles_per_app"]]
        apps.append({
            "app": app,
            "minutes": round(mins, 1),
            "duration": fmt_dur(mins),
            "titles": [{"title": t, "duration": fmt_dur(m)} for t, m in top],
        })

    return {
        "available": True,
        "first_sample": samples[0]["ts"],
        "last_sample": samples[-1]["ts"],
        "sample_count": len(samples),
        "active_duration": fmt_dur(sum(per_app.values())),
        "idle_duration": fmt_dur(idle_min),
        "apps": apps,
    }


# --------------------------------------------------------------------------- #
# collector: meetings attended (from Teams/Zoom window titles in the samples)
# --------------------------------------------------------------------------- #

_TEAMS_NAV = {"chat", "activity", "calendar", "calls", "files", "teams",
              "notifications", "apps", "copilot", "people", "search",
              "onedrive", "microsoft teams", ""}


def _meeting_title(app: str, title: str) -> str | None:
    """Meeting name from a window title, or None if it's not a meeting window."""
    if not title or title == "(excluded)":
        return None
    first = title.split(" | ")[0].strip()
    if app in ("MSTeams", "Microsoft Teams"):
        return None if first.lower() in _TEAMS_NAV else first
    if title.strip().lower() in {"zoom", "zoom workplace", "zoom.us", "webex"}:
        return None
    return first or None


def collect_meetings_attended(cfg: dict, day: date_cls) -> dict:
    acfg = cfg["activity"]
    if not acfg["enabled"]:
        return {"available": False, "reason": "activity tracking disabled"}
    apps = set(acfg.get("meeting_apps",
                        ["MSTeams", "Microsoft Teams", "zoom.us", "Zoom", "Webex"]))
    path = expand(cfg["state_dir"]) / "raw" / day.strftime("%Y-%m-%d") / "activity.jsonl"
    if not path.is_file():
        return {"available": False, "reason": "no activity samples for this day"}

    per_title, presence_samples = defaultdict(list), 0
    with open(path, errors="replace") as fh:
        for line in fh:
            line = line.strip()
            if not line:
                continue
            try:
                rec = json.loads(line)
                dt = datetime.fromisoformat(rec["ts"])
            except (json.JSONDecodeError, KeyError, ValueError):
                continue

            # a meeting window open at all means the call was live; samples written
            # before this field existed only know what was frontmost
            windows = rec.get("meeting_windows")
            if windows is None:
                if rec.get("app") in apps:
                    t = _meeting_title(rec["app"], rec.get("title", ""))
                    if t:
                        per_title[t].append(dt)
                continue

            presence_samples += 1
            seen = set()
            for wapp, wtitle in windows:
                if wapp not in apps or is_excluded_app(acfg, wapp):
                    continue
                t = _meeting_title(wapp, wtitle)
                if t and t not in seen:
                    seen.add(t)
                    per_title[t].append(dt)

    cap = float(acfg.get("sample_gap_cap_minutes", 2))
    meetings = []
    for title, times in per_title.items():
        for run in _group_runs(times):
            mins = _spent(run, cap)
            if mins < acfg["min_minutes_to_report"]:
                continue
            meetings.append(_meeting_entry(title, run, mins))
    meetings.sort(key=lambda m: m["start"])
    meetings, _g = drop_paused(pause_windows(cfg, day), meetings, "start", "end")
    return {"available": True, "count": len(meetings), "meetings": meetings,
            "basis": "window presence" if presence_samples else "frontmost window only"}


def _spent(times: list, cap: float) -> float:
    """Minutes covered by a run of samples, each gap capped so sleep can't inflate."""
    times.sort()
    mins = 1.0
    for a, b in zip(times, times[1:]):
        mins += min((b - a).total_seconds() / 60.0, cap)
    return mins


# A same-named meeting can occur twice in one day. Grouping only by title would
# report one entry spanning both, so a long silence starts a new occurrence.
# Short gaps stay inside one meeting: a dropped call or a reconnect is not a
# second meeting.
MEETING_SPLIT_GAP_MINUTES = 5.0


def _group_runs(times: list) -> list[list]:
    times = sorted(times)
    runs = [[times[0]]] if times else []
    for prev, cur in zip(times, times[1:]):
        if (cur - prev).total_seconds() / 60.0 > MEETING_SPLIT_GAP_MINUTES:
            runs.append([cur])
        else:
            runs[-1].append(cur)
    return runs


def _meeting_entry(title: str, times: list, mins: float) -> dict:
    return {"title": title,
            "start": times[0].isoformat(timespec="seconds"),
            "end": times[-1].isoformat(timespec="seconds"),
            "duration": fmt_dur(mins),
            "minutes": round(mins, 1)}


def collect_unjoined_meetings(cfg: dict, day: date_cls, attended=None,
                              calendar=None) -> dict:
    """Booked meetings that ran without being joined on this Mac.

    Best-effort and NOT a list of the day's calendar: Outlook's status menu only
    surfaces a meeting until it is joined, and some booked meetings never appear.
    Titles also present in meetings_attended are dropped, so this never claims a
    meeting was missed when the call was in fact joined.
    """
    scfg = cfg.get("unjoined_meetings") or {}
    if not scfg.get("enabled", True):
        return {"available": False, "reason": "disabled"}
    if (calendar or {}).get("available") and (calendar or {}).get("events"):
        return {"available": False,
                "reason": "superseded by the calendar feed for this day"}
    acfg = cfg["activity"]
    path = expand(cfg["state_dir"]) / "raw" / day.strftime("%Y-%m-%d") / "activity.jsonl"
    if not path.is_file():
        return {"available": False, "reason": "no activity samples for this day"}

    per_title, seen_field = defaultdict(list), False
    with open(path, errors="replace") as fh:
        for line in fh:
            line = line.strip()
            if not line:
                continue
            try:
                rec = json.loads(line)
                dt = datetime.fromisoformat(rec["ts"])
            except (json.JSONDecodeError, KeyError, ValueError):
                continue
            if "scheduled_now" not in rec:
                continue
            seen_field = True
            t = rec["scheduled_now"]
            if t and t != "(excluded)":
                per_title[t].append(dt)

    if not seen_field:
        return {"available": False,
                "reason": "no samples carry scheduled_now yet (sampler predates it)"}

    joined = {m["title"] for m in (attended or {}).get("meetings") or []}
    cap = float(acfg.get("sample_gap_cap_minutes", 2))
    meetings = [_meeting_entry(t, run, _spent(run, cap))
                for t, times in per_title.items() if t not in joined
                for run in _group_runs(times)]
    meetings.sort(key=lambda m: m["start"])
    return {"available": True, "count": len(meetings), "meetings": meetings,
            "source": "Outlook status menu",
            "caveat": "booked meetings seen running but not joined on this Mac; "
                      "not a complete calendar — Outlook stops surfacing a meeting "
                      "once it is joined and does not surface all of them"}


# --------------------------------------------------------------------------- #
# collector: git
# --------------------------------------------------------------------------- #

def _find_repos_under(root: Path, max_depth: int, excludes: list[str]) -> list[Path]:
    """find .git dirs under root, pruning heavy/irrelevant trees."""
    cmd = ["find", str(root), "-maxdepth", str(max_depth)]
    if excludes:
        cmd.append("(")
        for i, name in enumerate(excludes):
            if i:
                cmd.append("-o")
            cmd += ["-name", name]
        cmd += [")", "-prune", "-o"]
    cmd += ["-type", "d", "-name", ".git", "-print"]
    rc, out, _ = run(cmd, timeout=180)
    if rc not in (0, 1):
        return []
    return [Path(line).parent for line in out.splitlines() if line.strip()]


DEFAULT_EXCLUDES = ["node_modules", "Library", ".Trash", "Applications", ".venv",
                    "venv", ".tox", "vendor", "Pods", ".cache", "CloudStorage",
                    "Music", "Movies", "Pictures", "site-packages"]


def find_git_repos(cfg: dict) -> list[Path]:
    gcfg = cfg["git"]
    cache_path = expand(cfg["state_dir"]) / "repo-cache.json"
    ttl = gcfg.get("repo_cache_hours", 24) * 3600
    if cache_path.is_file() and (time.time() - cache_path.stat().st_mtime) < ttl:
        try:
            with open(cache_path) as fh:
                return [Path(p) for p in json.load(fh)]
        except (json.JSONDecodeError, OSError):
            pass

    excludes = gcfg.get("exclude_dir_names", DEFAULT_EXCLUDES)
    repos: list[Path] = []
    for root in gcfg["scan_roots"]:
        rp = expand(root)
        if rp.is_dir():
            repos.extend(_find_repos_under(rp, gcfg["scan_max_depth"], excludes))

    repos_s = sorted({str(r) for r in repos})
    try:
        cache_path.parent.mkdir(parents=True, exist_ok=True)
        with open(cache_path, "w") as fh:
            json.dump(repos_s, fh, indent=2)
    except OSError:
        pass
    return [Path(r) for r in repos_s]


def git_author_emails(cfg: dict) -> list[str]:
    emails = set(cfg["git"].get("extra_author_emails", []))
    rc, out, _ = run(["git", "config", "--get-all", "user.email"], timeout=10)
    if rc == 0:
        emails.update(e.strip() for e in out.splitlines() if e.strip())
    return sorted(emails)


def _commit_detail(repo: Path, sha: str, redact) -> dict:
    """Body + numstat for one commit. Best-effort; empty dict on failure."""
    d = {}
    rc, out, _ = run(["git", "-C", str(repo), "show", sha, "--numstat",
                      "--pretty=format:"], timeout=15)
    if rc == 0:
        files, ins, dels = [], 0, 0
        for line in out.splitlines():
            m = re.match(r"^(\d+|-)\t(\d+|-)\t(.+)$", line)
            if m:
                ins += 0 if m.group(1) == "-" else int(m.group(1))
                dels += 0 if m.group(2) == "-" else int(m.group(2))
                files.append(m.group(3))
        d = {"files_changed": len(files), "insertions": ins, "deletions": dels,
             "files": files[:8]}
    rc, body, _ = run(["git", "-C", str(repo), "show", "-s",
                       "--pretty=format:%b", sha], timeout=15)
    if rc == 0 and body.strip():
        d["body"] = redact(body.strip())[:300]
    return d


def collect_git(cfg: dict, day: date_cls, bounds=None) -> dict:
    if not cfg["git"]["enabled"]:
        return {"available": False, "reason": "disabled"}
    start, end = bounds or window_bounds(cfg, day)
    since, until = start.isoformat(), end.isoformat()
    emails = git_author_emails(cfg)
    repos = find_git_repos(cfg)
    if not repos:
        return {"available": False, "reason": "no git repos found under scan_roots"}

    redact = Redactor(cfg)
    _pw = pause_windows(cfg, day)
    results = []
    for repo in repos:
        # repos may set a local user.email different from the global one;
        # filter on the repo's effective identity too
        repo_emails = set(emails)
        rc0, eff, _ = run(["git", "-C", str(repo), "config", "user.email"],
                          timeout=10)
        if rc0 == 0 and eff.strip():
            repo_emails.add(eff.strip())
        args = ["git", "-C", str(repo), "log", "--all", "--no-merges",
                f"--since={since}", f"--until={until}",
                "--pretty=format:%h\x1f%an\x1f%ae\x1f%aI\x1f%s"]
        for em in sorted(repo_emails):
            args.append(f"--author={em}")
        rc, out, _ = run(args, timeout=45)
        commits = []
        if rc == 0 and out:
            for line in out.splitlines():
                f = line.split("\x1f")
                if len(f) == 5:
                    c = {"sha": f[0], "author": f[1], "email": f[2],
                         "at": f[3], "subject": redact(f[4])}
                    c.update(_commit_detail(repo, f[0], redact))
                    commits.append(c)

        branches = set()
        rc2, out2, _ = run(
            ["git", "-C", str(repo), "reflog", "--date=iso-strict",
             f"--since={since}", f"--until={until}",
             "--pretty=format:%gs"], timeout=30)
        if rc2 == 0 and out2:
            for line in out2.splitlines():
                m = re.search(r"checkout: moving from \S+ to (\S+)", line)
                if m:
                    branches.add(m.group(1))

        commits, gone = drop_paused(_pw, commits, "at")
        if gone:
            log_line(cfg, f"pause: hid {gone} commit(s) in {repo.name} on {day}")
        if commits or branches:
            rc3, cur, _ = run(["git", "-C", str(repo), "rev-parse",
                               "--abbrev-ref", "HEAD"], timeout=10)
            results.append({
                "repo": repo.name,
                "path": str(repo),
                "current_branch": cur if rc3 == 0 else "",
                "commits": commits,
                "branches_touched": sorted(branches),
            })

    return {
        "available": True,
        "author_emails": emails,
        "repos_scanned": len(repos),
        "repos_with_activity": results,
        "total_commits": sum(len(r["commits"]) for r in results),
    }


# --------------------------------------------------------------------------- #
# collector: shell history
# --------------------------------------------------------------------------- #

_ZSH_EXT = re.compile(r"^:\s*(\d+):(\d+);(.*)$")


def collect_shell(cfg: dict, day: date_cls, bounds=None) -> dict:
    scfg = cfg["shell"]
    if not scfg["enabled"]:
        return {"available": False, "reason": "disabled"}
    start, end = bounds or window_bounds(cfg, day)
    lo, hi = start.timestamp(), end.timestamp()
    redact = Redactor(cfg)
    ignores = []
    for pat in scfg["ignore_patterns"]:
        try:
            ignores.append(re.compile(pat))
        except re.error:
            pass

    entries, timestamped_any, files_read = [], False, []
    for hf in scfg["history_files"]:
        path = expand(hf)
        if not path.is_file():
            continue
        files_read.append(str(path))
        with open(path, "rb") as fh:
            raw = fh.read().decode("utf-8", errors="replace")

        pending_ts, buf = None, None
        for line in raw.splitlines():
            m = _ZSH_EXT.match(line)
            if m:
                if buf is not None and pending_ts is not None:
                    entries.append((pending_ts, buf))
                timestamped_any = True
                pending_ts, buf = int(m.group(1)), m.group(3)
            elif buf is not None:
                buf += " " + line.strip()
        if buf is not None and pending_ts is not None:
            entries.append((pending_ts, buf))

    if not files_read:
        return {"available": False, "reason": "no history files found"}
    if not timestamped_any:
        return {"available": False,
                "reason": "history has no timestamps — enable EXTENDED_HISTORY in ~/.zshrc"}

    cmds, seen = [], set()
    for ts, cmd in sorted(entries):
        if not (lo <= ts <= hi):
            continue
        cmd = cmd.strip().rstrip("\\").strip()
        if not cmd or any(rx.search(cmd) for rx in ignores):
            continue
        key = cmd[:160]
        if key in seen:
            continue
        seen.add(key)
        cmds.append({"at": datetime.fromtimestamp(ts).isoformat(timespec="seconds"),
                     "cmd": redact(cmd)[:400]})

    cmds, _g = drop_paused(pause_windows(cfg, day), cmds, "at")
    cmds = cmds[: scfg["max_commands"]]
    return {"available": True, "files": files_read, "count": len(cmds), "commands": cmds}


# --------------------------------------------------------------------------- #
# collector: calendar
# --------------------------------------------------------------------------- #

def _calendar_via_icalbuddy(day: date_cls, timeout: int) -> list[dict] | None:
    if not shutil.which("icalBuddy"):
        return None
    # -ea: skip all-day events (holiday/birthday calendars), we want meetings
    rc, out, _ = run(
        ["icalBuddy", "-ea", "-nc", "-nrd", "-df", "%Y-%m-%d", "-tf", "%H:%M",
         "-ps", "|;|", "-b", "", "eventsFrom:" + day.isoformat(),
         "to:" + day.isoformat()],
        timeout=timeout)
    if rc != 0:
        return None
    events = []
    for block in out.split("\n"):
        block = block.strip()
        if block:
            events.append({"raw": block})
    return events or []


_OUTLOOK_SCRIPT = '''
set d1 to (current date)
set day of d1 to {day}
set month of d1 to {month}
set year of d1 to {year}
set time of d1 to 0
set d2 to d1 + (1 * days)
set out to ""
tell application "Microsoft Outlook"
  repeat with c in every calendar
    try
      set evs to (every calendar event of c whose start time is greater than or equal to d1 and start time is less than d2)
      repeat with e in evs
        try
          if (all day flag of e) is false then
            set fb to ""
            try
              set fb to free busy status of e as string
            end try
            set out to out & (subject of e) & tab & ((start time of e) as string) & tab & ((end time of e) as string) & tab & fb & linefeed
          end if
        end try
      end repeat
    end try
  end repeat
end tell
return out
'''


def _new_outlook_active() -> bool:
    rc, out, _ = run(["defaults", "read", "com.microsoft.Outlook",
                      "IsRunningNewOutlook"], timeout=10)
    return rc == 0 and out.strip() == "1"

_CALENDAR_APP_SCRIPT = '''
set d1 to (current date)
set day of d1 to {day}
set month of d1 to {month}
set year of d1 to {year}
set time of d1 to 0
set d2 to d1 + (1 * days)
set out to ""
tell application "Calendar"
  repeat with c in calendars
    try
      set evs to (every event of c whose start date is greater than or equal to d1 and start date is less than d2)
      repeat with e in evs
        set out to out & (summary of e) & tab & ((start date of e) as string) & tab & ((end date of e) as string) & linefeed
      end repeat
    end try
  end repeat
end tell
return out
'''


# --------------------------------------------------------------------------- #
# calendar from an exported / published .ics — the only real calendar source
#
# Outlook writes Windows zone names and leans on RRULE, so occurrences have to be
# expanded here. RECURRENCE-ID overrides matter: a rescheduled occurrence moves,
# which is why the calendar can legitimately disagree with what was attended.
# --------------------------------------------------------------------------- #

_WINDOWS_TZ = {
    "India Standard Time": "Asia/Kolkata",
    "Eastern Standard Time": "America/Toronto",
    "Pacific Standard Time": "America/Los_Angeles",
    "Turkey Standard Time": "Europe/Istanbul",
    "GMT Standard Time": "Europe/London",
    "W. Europe Standard Time": "Europe/Berlin",
    "Central Standard Time": "America/Chicago",
    "Mountain Standard Time": "America/Denver",
    "Tokyo Standard Time": "Asia/Tokyo",
    "China Standard Time": "Asia/Shanghai",
    "AUS Eastern Standard Time": "Australia/Sydney",
    "UTC": "UTC",
}
_BYDAY = {"MO": 0, "TU": 1, "WE": 2, "TH": 3, "FR": 4, "SA": 5, "SU": 6}
_CANCELLED_RE = re.compile(r"^\s*cancell?ed:\s*", re.I)
_ICS_CACHE: dict = {}


def _ics_unfold(text: str) -> list[str]:
    out, buf = [], ""
    for ln in text.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        if ln[:1] in (" ", "\t"):
            buf += ln[1:]
        else:
            if buf:
                out.append(buf)
            buf = ln
    if buf:
        out.append(buf)
    return out


def _ics_zone(tzid: str | None):
    if not tzid:
        return None
    try:
        return ZoneInfo(_WINDOWS_TZ.get(tzid, tzid))
    except Exception:
        return None


def _ics_dt(value: str, params: dict) -> tuple[datetime, bool]:
    v = value.strip()
    if params.get("VALUE") == "DATE" or re.fullmatch(r"\d{8}", v):
        return datetime.strptime(v[:8], "%Y%m%d"), True
    if v.endswith("Z"):
        return datetime.strptime(v[:15], "%Y%m%dT%H%M%S").replace(
            tzinfo=ZoneInfo("UTC")), False
    dt = datetime.strptime(v[:15], "%Y%m%dT%H%M%S")
    z = _ics_zone(params.get("TZID"))
    return (dt.replace(tzinfo=z) if z else dt.astimezone()), False


def _ics_line(line: str) -> tuple[str, dict, str]:
    head, _, value = line.partition(":")
    parts = head.split(";")
    params = {}
    for p in parts[1:]:
        k, _, v = p.partition("=")
        params[k.upper()] = v
    return parts[0].upper(), params, value


def _ics_read(text: str) -> list[dict]:
    """VEVENT blocks only — VTIMEZONE carries RRULEs that are not event rules."""
    events, cur, in_tz = [], None, 0
    for line in _ics_unfold(text):
        name, params, value = _ics_line(line)
        v = value.strip().upper()
        if name == "BEGIN" and v == "VTIMEZONE":
            in_tz += 1
            continue
        if name == "END" and v == "VTIMEZONE":
            in_tz -= 1
            continue
        if in_tz:
            continue
        if name == "BEGIN" and v == "VEVENT":
            cur = {"exdates": []}
            continue
        if name == "END" and v == "VEVENT":
            if cur is not None:
                events.append(cur)
            cur = None
            continue
        if cur is None:
            continue
        try:
            if name == "SUMMARY":
                cur["summary"] = value.strip()
            elif name == "LOCATION":
                cur["location"] = value.strip()
            elif name == "UID":
                cur["uid"] = value.strip()
            elif name == "DTSTART":
                cur["start"], cur["all_day"] = _ics_dt(value, params)
            elif name == "DTEND":
                cur["end"], _ = _ics_dt(value, params)
            elif name == "RECURRENCE-ID":
                cur["recurrence_id"], _ = _ics_dt(value, params)
            elif name == "RRULE":
                cur["rrule"] = {k.upper(): val for k, _, val in
                                (p.partition("=") for p in value.strip().split(";")) if k}
            elif name == "EXDATE":
                for piece in value.split(","):
                    cur["exdates"].append(_ics_dt(piece, params)[0])
            elif name == "STATUS":
                cur["status"] = value.strip().upper()
            elif name == "TRANSP":
                cur["transp"] = value.strip().upper()
            elif name == "X-MICROSOFT-CDO-BUSYSTATUS":
                cur["busy"] = value.strip().upper()
            elif name == "X-MICROSOFT-CDO-ALLDAYEVENT":
                cur["all_day"] = cur.get("all_day") or v == "TRUE"
        except ValueError:
            continue
    return events


def _ics_expand(ev: dict, lo: date_cls, hi: date_cls) -> list[datetime]:
    start = ev["start"]
    rr = ev.get("rrule")
    if not rr:
        return [start] if lo <= start.date() <= hi else []
    freq = rr.get("FREQ", "").upper()
    interval = max(int(rr.get("INTERVAL") or 1), 1)
    count = int(rr["COUNT"]) if rr.get("COUNT") else None
    until = None
    if rr.get("UNTIL"):
        try:
            until = _ics_dt(rr["UNTIL"], {})[0]
        except ValueError:
            until = None
    days = [_BYDAY[d] for d in rr.get("BYDAY", "").split(",") if d in _BYDAY]

    if freq == "WEEKLY":
        wanted = sorted(days or [start.weekday()])
        cursor = start.date() - timedelta(days=start.weekday())
        out, n, limit = [], 0, hi + timedelta(days=7)
        while cursor <= limit:
            for wd in wanted:
                d = cursor + timedelta(days=wd)
                if d < start.date():
                    continue
                occ = start.replace(year=d.year, month=d.month, day=d.day)
                if until and occ > until:
                    return out
                n += 1
                if count and n > count:
                    return out
                if lo <= d <= hi:
                    out.append(occ)
            cursor += timedelta(weeks=interval)
        return out

    if freq == "DAILY":
        out, n, d = [], 0, start
        while d.date() <= hi:
            if until and d > until:
                break
            n += 1
            if count and n > count:
                break
            if d.date() >= lo:
                out.append(d)
            d += timedelta(days=interval)
        return out

    return [start] if lo <= start.date() <= hi else []


def _ics_entry(ev: dict, start: datetime, end: datetime) -> dict:
    return {
        "subject": (ev.get("summary") or "").strip(),
        "start": start.astimezone().replace(tzinfo=None).isoformat(timespec="seconds"),
        "end": end.astimezone().replace(tzinfo=None).isoformat(timespec="seconds"),
        "location": (ev.get("location") or "").strip(),
        "busy": ev.get("busy") or "BUSY",
        "free": (ev.get("transp") or "") == "TRANSPARENT",
        "all_day": bool(ev.get("all_day")),
    }


def ics_occurrences(text: str, lo: date_cls, hi: date_cls) -> list[dict]:
    events = _ics_read(text)
    masters = [e for e in events if "recurrence_id" not in e and e.get("start")]
    overrides = {(e.get("uid"), e["recurrence_id"]): e
                 for e in events if "recurrence_id" in e}
    out = []
    for m in masters:
        dur = (m["end"] - m["start"]) if m.get("end") else timedelta(minutes=30)
        for occ in _ics_expand(m, lo, hi):
            if any(abs((occ - x).total_seconds()) < 60 for x in m["exdates"]):
                continue
            ov = overrides.get((m.get("uid"), occ))
            src = ov or m
            summary = (src.get("summary") or m.get("summary") or "").strip()
            if _CANCELLED_RE.match(summary) or src.get("status") == "CANCELLED":
                continue
            s = src["start"] if ov else occ
            # a plain occurrence ends dur after ITS start; only an override carries
            # its own DTEND. Reusing the master's DTEND would stamp the series'
            # original date on every occurrence.
            e = (ov.get("end") or (s + dur)) if ov else (s + dur)
            merged = {**m, **src, "summary": summary}
            out.append(_ics_entry(merged, s, e))
    seen = {(o["subject"], o["start"]) for o in out}
    for ov in overrides.values():          # an override can land outside the series
        s = ov.get("start")
        summary = (ov.get("summary") or "").strip()
        if not s or not (lo <= s.astimezone().date() <= hi):
            continue
        if _CANCELLED_RE.match(summary) or ov.get("status") == "CANCELLED":
            continue
        e = ov.get("end") or (s + timedelta(minutes=30))
        entry = _ics_entry(ov, s, e)
        if (entry["subject"], entry["start"]) not in seen:
            out.append(entry)
    # An override can move an occurrence to another date, so the effective start
    # has to be re-checked: expanding a single day must not return the day it moved to.
    out = [o for o in out
           if lo <= date_cls.fromisoformat(o["start"][:10]) <= hi]
    out.sort(key=lambda x: x["start"])
    return out


def _ics_text(cfg: dict) -> tuple[str | None, str]:
    """ICS body plus where it came from. A published URL beats a stale export."""
    ccfg = cfg["calendar"]
    url_file = ccfg.get("ics_url_file")
    if url_file:
        p = expand(url_file)
        if p.is_file():
            url = p.read_text().strip()
            if url:
                key = ("url", url)
                if key in _ICS_CACHE:
                    return _ICS_CACHE[key], "published URL"
                try:
                    import urllib.request
                    with urllib.request.urlopen(
                            url, timeout=ccfg.get("fetch_timeout_seconds", 30)) as r:
                        body = r.read().decode("utf-8", "replace")
                    if "BEGIN:VCALENDAR" in body:
                        _ICS_CACHE[key] = body
                        return body, "published URL"
                    log_line(cfg, "calendar: URL returned no VCALENDAR")
                except Exception as exc:
                    log_line(cfg, f"calendar: URL fetch failed {exc!r}, trying local file")
    path = ccfg.get("ics_path")
    if path:
        p = expand(path)
        if p.is_file():
            key = ("file", str(p), p.stat().st_mtime)
            if key not in _ICS_CACHE:
                _ICS_CACHE[key] = p.read_text(errors="replace")
            return _ICS_CACHE[key], f"exported file ({p.name})"
    return None, "none"


def collect_calendar_ics(cfg: dict, day: date_cls) -> dict | None:
    text, origin = _ics_text(cfg)
    if not text:
        return None
    ccfg = cfg["calendar"]
    try:
        occ = ics_occurrences(text, day, day)
    except Exception as exc:
        log_line(cfg, f"calendar: ICS parse failed {exc!r}")
        return None
    events, skipped = [], 0
    for o in occ:
        if o["all_day"] and not ccfg.get("include_all_day", False):
            skipped += 1
            continue
        if o["free"] and not ccfg.get("include_free", False):
            skipped += 1
            continue
        if o["busy"] == "OOF" and not ccfg.get("include_out_of_office", False):
            o = {**o, "out_of_office": True}
        events.append(o)
    events, gone = drop_paused(pause_windows(cfg, day), events, "start", "end")
    if gone:
        log_line(cfg, f"pause: hid {gone} calendar event(s) on {day}")
    return {"available": True, "source": f"ICS via {origin}",
            "events": events, "count": len(events), "skipped": skipped}


def collect_calendar(cfg: dict, day: date_cls) -> dict:
    ccfg = cfg["calendar"]
    if not ccfg["enabled"]:
        return {"available": False, "reason": "disabled"}
    ics = collect_calendar_ics(cfg, day)
    if ics is not None:
        return ics
    timeout = ccfg["timeout_seconds"]
    prefer = ccfg.get("prefer", "auto")

    # an empty result from one source falls through to the next; only a
    # non-empty result short-circuits, so Outlook still gets asked when the
    # Apple calendar has no meetings on it
    empty_sources = []

    if prefer in ("auto", "icalbuddy"):
        ib = _calendar_via_icalbuddy(day, timeout)
        if ib:
            return {"available": True, "source": "icalBuddy", "events": ib}
        if ib is not None:
            empty_sources.append("icalBuddy")

    subs = {"day": day.day, "month": day.month, "year": day.year,
            "idx": ccfg["outlook_calendar_index"]}
    attempts = []
    if prefer in ("auto", "outlook"):
        attempts.append(("Microsoft Outlook", _OUTLOOK_SCRIPT))
    if prefer in ("auto", "calendar"):
        attempts.append(("Calendar.app", _CALENDAR_APP_SCRIPT))

    errors = []
    for name, tmpl in attempts:
        rc, out, err = osascript(tmpl.format(**subs), timeout=timeout)
        if rc != 0:
            errors.append(f"{name}: {err[:160] or 'rc=' + str(rc)}")
            continue
        events = []
        for line in out.splitlines():
            f = line.split("\t")
            if f and f[0].strip():
                ev = {"subject": f[0].strip(),
                      "start": f[1].strip() if len(f) > 1 else "",
                      "end": f[2].strip() if len(f) > 2 else ""}
                status = f[3].strip() if len(f) > 3 else ""
                if status and status not in ("busy",):
                    ev["status"] = status  # tentative / out of office / free
                events.append(ev)
        if events:
            return {"available": True, "source": name, "events": events}
        empty_sources.append(name)

    if errors:
        return {"available": False, "reason": "; ".join(errors)}
    if "Microsoft Outlook" in empty_sources and _new_outlook_active():
        return {"available": False, "reason":
                "Outlook runs in 'New Outlook' mode — AppleScript cannot read its "
                "events. Add the Exchange account in System Settings → Internet "
                "Accounts (Calendars on) so Calendar.app/icalBuddy can see meetings"}
    if empty_sources:
        return {"available": True, "source": " / ".join(empty_sources), "events": []}
    return {"available": False, "reason": "no calendar source"}


# --------------------------------------------------------------------------- #
# collector: Claude Code sessions
# --------------------------------------------------------------------------- #

def _extract_text(content) -> str:
    if isinstance(content, str):
        return content
    if isinstance(content, list):
        out = []
        for blk in content:
            if isinstance(blk, dict) and blk.get("type") == "text":
                out.append(str(blk.get("text", "")))
            elif isinstance(blk, str):
                out.append(blk)
        return " ".join(out)
    if isinstance(content, dict):
        return _extract_text(content.get("content", ""))
    return ""


def _tool_names(content) -> list[str]:
    names = []
    if isinstance(content, list):
        for blk in content:
            if isinstance(blk, dict) and blk.get("type") == "tool_use":
                n = blk.get("name")
                if n:
                    names.append(str(n))
    return names


def _tool_calls(content):
    """(name, input) for each tool call — the inputs say what was actually done."""
    if isinstance(content, list):
        for blk in content:
            if isinstance(blk, dict) and blk.get("type") == "tool_use" and blk.get("name"):
                yield str(blk["name"]), (blk.get("input") or {})


_TEST_HINT = re.compile(r"(?i)(^|[/_.-])(test|tests|spec|specs|__tests__)([/_.-]|$)")

# Not everything recorded as a "user" turn was typed by the user: the harness injects
# tool scaffolding, slash-command echoes and interruption markers. Logging those as
# intent is worse than logging nothing.
_PROMPT_NOISE = re.compile(
    r"^(?:\[Request interrupted"
    r"|Repo worktree:"
    r"|Base directory for this skill:"
    r"|Caveat: The messages below"
    r"|This session is being continued"
    r"|<)", re.I)


def _is_real_prompt(txt: str) -> bool:
    return bool(txt) and not _PROMPT_NOISE.match(txt.strip())


def _top_n(items: list[str], n: int, keep_order: bool = False) -> list[str]:
    """De-duplicate, keeping either first-seen order or the most frequent first."""
    if keep_order:
        seen, out = set(), []
        for x in items:
            if x not in seen:
                seen.add(x)
                out.append(x)
        return out[:n]
    counts = defaultdict(int)
    for x in items:
        counts[x] += 1
    return [x for x, _ in sorted(counts.items(), key=lambda kv: (-kv[1], kv[0]))][:n]


def _short_path(p: str) -> str:
    """Repo-relative where possible: informative without pasting absolute paths."""
    s = str(p or "")
    for marker in ("/Documents/", "/scratchpad/"):
        i = s.find(marker)
        if i >= 0:
            return s[i + len(marker):]
    return Path(s).name or s


def collect_claude_code(cfg: dict, day: date_cls, bounds=None) -> dict:
    kcfg = cfg["claude_code"]
    if not kcfg["enabled"]:
        return {"available": False, "reason": "disabled"}
    root = expand(kcfg["projects_dir"])
    if not root.is_dir():
        return {"available": False, "reason": f"{root} not found"}

    start, end = bounds or window_bounds(cfg, day)
    redact = Redactor(cfg)
    sessions = []

    for jf in sorted(root.rglob("*.jsonl")):
        try:
            mtime = datetime.fromtimestamp(jf.stat().st_mtime)
        except OSError:
            continue
        if mtime < start - timedelta(days=1):
            continue

        prompts, tools, cwds, times, turns = [], defaultdict(int), set(), [], 0
        is_self = False
        titles, branches, edited, commands, skills = {}, set(), [], [], []
        kinds, sidechain = set(), False
        try:
            with open(jf, errors="replace") as fh:
                for line in fh:
                    line = line.strip()
                    if not line:
                        continue
                    try:
                        rec = json.loads(line)
                    except json.JSONDecodeError:
                        continue
                    if not isinstance(rec, dict):
                        continue
                    # title and last-prompt records carry no timestamp, so they have
                    # to be read before the in-window filter or they are lost
                    rtype = rec.get("type")
                    if rtype in ("custom-title", "ai-title"):
                        key = "custom" if rtype == "custom-title" else "ai"
                        val = rec.get("customTitle") or rec.get("aiTitle")
                        if val:
                            titles[key] = str(val)[:200]
                        continue
                    # `last-prompt` is deliberately ignored: it is file-global, so on a
                    # session spanning days it reports a prompt from another day. The
                    # last in-window prompt is the honest "where this was left".
                    ts = rec.get("timestamp") or rec.get("ts") or ""
                    try:
                        dt = datetime.fromisoformat(str(ts).replace("Z", "+00:00"))
                        dt = dt.astimezone().replace(tzinfo=None) if dt.tzinfo else dt
                    except (ValueError, TypeError):
                        continue
                    if not (start <= dt <= end):
                        continue
                    times.append(dt)
                    if rec.get("cwd"):
                        cwds.add(str(rec["cwd"]))
                    if rec.get("gitBranch"):
                        branches.add(str(rec["gitBranch"]))
                    if rec.get("isSidechain"):
                        sidechain = True
                    if rec.get("sessionKind"):
                        kinds.add(str(rec["sessionKind"]))
                    msg = rec.get("message") or {}
                    role = rec.get("type") or (msg.get("role") if isinstance(msg, dict) else "")
                    content = msg.get("content") if isinstance(msg, dict) else None
                    if role == "user":
                        txt = _extract_text(content).strip()
                        if is_self_prompt(txt):
                            is_self = True
                            break
                        if _is_real_prompt(txt):
                            # collapse newlines: a multi-line prompt would otherwise
                            # put continuation text at column 0 and fake a heading
                            flat = " ".join(redact(txt).split())
                            prompts.append(flat[: kcfg["prompt_truncate_chars"]])
                    elif role == "assistant":
                        turns += 1
                        for name, inp in _tool_calls(content):
                            tools[name] += 1
                            if name in ("Edit", "Write", "NotebookEdit"):
                                if inp.get("file_path"):
                                    edited.append(_short_path(inp["file_path"]))
                            elif name in ("Bash", "Monitor"):
                                # the description is the human-readable intent of the
                                # command, which is what a work log actually wants
                                d = str(inp.get("description") or "").strip()
                                if d:
                                    commands.append(redact(" ".join(d.split()))[:120])
                            elif name == "Skill" and inp.get("skill"):
                                skills.append(str(inp["skill"])[:60])
        except OSError:
            continue

        if is_self:
            continue
        if not times:
            continue
        project = sorted(cwds)[0] if cwds else jf.parent.name
        if project in (kcfg.get("exclude_projects") or []):
            continue
        times.sort()
        # active time = sum of gaps between entries, capped at 5 min so a session
        # left open over lunch doesn't read as three hours of work
        active = 0.0
        for a, b in zip(times, times[1:]):
            active += min((b - a).total_seconds() / 60.0, 5.0)
        sessions.append({
            "session_file": jf.name,
            "project": project,
            "start": min(times).isoformat(timespec="seconds"),
            "end": max(times).isoformat(timespec="seconds"),
            "span": fmt_dur((max(times) - min(times)).total_seconds() / 60),
            "duration": fmt_dur(active),
            "assistant_turns": turns,
            "title": titles.get("custom") or titles.get("ai") or "",
            "ai_title": titles.get("ai", ""),
            "branches": sorted(branches),
            "files_changed": _top_n(edited, kcfg.get("max_files_per_session", 12)),
            "files_changed_total": len(set(edited)),
            "tests_touched": sorted({f for f in edited if _TEST_HINT.search(f)})[:8],
            "did": _top_n(commands, kcfg.get("max_commands_per_session", 14), keep_order=True),
            "command_count": len(commands),
            "skills_used": sorted(set(skills)),
            "subagent": sidechain or "bg" in kinds,
            "session_kind": sorted(kinds),
            "left_off": prompts[-1] if prompts else "",
            "tools_used": dict(sorted(tools.items(), key=lambda kv: -kv[1])),
            "prompts": prompts[: kcfg["max_prompts_per_session"]],
        })

    sessions.sort(key=lambda s: s["start"])
    sessions, _gone = drop_paused(pause_windows(cfg, day), sessions, "start", "end")
    if _gone:
        log_line(cfg, f"pause: hid {_gone} Claude Code session(s) on {day}")
    sessions, forks = _dedupe_forked_sessions(cfg, sessions)
    out = {"available": True, "session_count": len(sessions), "sessions": sessions}
    if forks:
        out["forked_duplicates_dropped"] = forks
    return out


def _dedupe_forked_sessions(cfg: dict, sessions: list[dict]) -> tuple[list[dict], int]:
    """Resuming a session copies the whole transcript into a new file.

    Both files then report the same start, end and turn count, so the day's duration
    and turns were counted twice. Identical transcripts collapse into one entry, with
    the other file's title kept — a fork is often relabelled and the label is useful.
    """
    kept: dict[tuple, dict] = {}
    dropped = 0
    for s in sessions:
        key = (s["start"], s["end"], s["assistant_turns"], s["project"])
        first = kept.get(key)
        if first is None:
            kept[key] = s
            continue
        dropped += 1
        for t in (s.get("title"), s.get("ai_title")):
            if t and t != first.get("title") and t not in first.setdefault("also_titled", []):
                first["also_titled"].append(t)
    if dropped:
        log_line(cfg, f"claude_code: collapsed {dropped} forked session copy(ies) "
                      f"that would have double-counted time")
    return list(kept.values()), dropped


# --------------------------------------------------------------------------- #
# collector: OneDrive / SharePoint local sync
# --------------------------------------------------------------------------- #

def collect_cloud(cfg: dict, day: date_cls, bounds=None) -> dict:
    ccfg = cfg["cloud"]
    if not ccfg["enabled"]:
        return {"available": False, "reason": "disabled"}
    start, end = bounds or window_bounds(cfg, day)
    lo, hi = start.timestamp(), end.timestamp()

    roots = [expand(r) for r in ccfg["scan_roots"]]
    roots = [r for r in roots if r.is_dir()]
    # also catch legacy "~/OneDrive - Org" style folders
    home = Path.home()
    if home.is_dir():
        for child in home.iterdir():
            if child.is_dir() and child.name.startswith(("OneDrive", "SharePoint")):
                roots.append(child)
    if not roots:
        return {"available": False,
                "reason": "no OneDrive/SharePoint sync folder found under ~/Library/CloudStorage"}

    ignores = []
    for pat in ccfg["ignore_patterns"]:
        try:
            ignores.append(re.compile(pat))
        except re.error:
            pass

    hits = []
    for root in roots:
        rc, out, _ = run(
            ["find", str(root), "-maxdepth", str(ccfg["scan_max_depth"]),
             "-type", "f", "-newermt", start.isoformat(),
             "!", "-newermt", end.isoformat()],
            timeout=180)
        if rc != 0:
            continue
        for line in out.splitlines():
            p = Path(line)
            if p.name in ccfg["ignore_names"]:
                continue
            # test both the bare filename (for patterns like ^~\$) and the full path
            if any(rx.search(p.name) or rx.search(str(p)) for rx in ignores):
                continue
            try:
                st = p.stat()
            except OSError:
                continue
            if not (lo <= st.st_mtime <= hi):
                continue
            try:
                rel = str(p.relative_to(root.parent))
            except ValueError:
                rel = str(p)
            hits.append({
                "file": rel,
                "modified": datetime.fromtimestamp(st.st_mtime).isoformat(timespec="seconds"),
                "kb": round(st.st_size / 1024, 1),
            })

    hits.sort(key=lambda h: h["modified"])
    hits, _g = drop_paused(pause_windows(cfg, day), hits, "modified")
    truncated = len(hits) > ccfg["max_files"]
    return {"available": True, "roots": [str(r) for r in roots],
            "count": len(hits), "truncated": truncated,
            "files": hits[: ccfg["max_files"]]}


# --------------------------------------------------------------------------- #
# collector: claude.ai export (manual drop)
# --------------------------------------------------------------------------- #

def collect_claude_export(cfg: dict, day: date_cls) -> dict:
    ecfg = cfg["claude_export"]
    if not ecfg["enabled"]:
        return {"available": False, "reason": "disabled"}
    inbox = expand(ecfg["inbox_dir"])
    if not inbox.is_dir():
        return {"available": False, "reason": f"no export dropped in {inbox}"}
    files = sorted(inbox.glob("*.json"))
    if not files:
        return {"available": False,
                "reason": f"no *.json in {inbox} (export from claude.ai → Settings → Export data)"}

    start, end = window_bounds(cfg, day)
    redact = Redactor(cfg)
    convos = []
    for f in files:
        try:
            with open(f, errors="replace") as fh:
                data = json.load(fh)
        except (json.JSONDecodeError, OSError):
            continue
        if isinstance(data, dict):
            data = data.get("conversations", [])
        if not isinstance(data, list):
            continue
        for c in data:
            if not isinstance(c, dict):
                continue
            ts = c.get("updated_at") or c.get("created_at") or ""
            try:
                dt = datetime.fromisoformat(str(ts).replace("Z", "+00:00"))
                dt = dt.astimezone().replace(tzinfo=None) if dt.tzinfo else dt
            except (ValueError, TypeError):
                continue
            if not (start <= dt <= end):
                continue
            msgs = c.get("chat_messages") or c.get("messages") or []
            first_user = ""
            for m in msgs:
                if isinstance(m, dict) and m.get("sender") in ("human", "user"):
                    first_user = _extract_text(m.get("content") or m.get("text") or "")
                    if first_user:
                        break
            convos.append({
                "name": redact(str(c.get("name") or "(untitled)"))[:160],
                "updated": dt.isoformat(timespec="seconds"),
                "message_count": len(msgs),
                "opening_prompt": redact(first_user)[:220],
            })

    convos.sort(key=lambda c: c["updated"])
    return {"available": True, "source_files": [f.name for f in files],
            "count": len(convos), "conversations": convos[: ecfg["max_conversations"]]}


# --------------------------------------------------------------------------- #
# digest
# --------------------------------------------------------------------------- #

# --------------------------------------------------------------------------- #
# collector: the team's MIS board export — tasks the user actually logged
# --------------------------------------------------------------------------- #

_MIS_CACHE: dict = {}


def _mis_rows_for(cfg: dict) -> tuple[list[dict], str] | tuple[None, str]:
    mcfg = cfg.get("mis_board") or {}
    p = expand(mcfg.get("path", "~/.worklog/mis-board.xlsx"))
    if not p.is_file():
        return None, f"{p} not found"
    key = (str(p), p.stat().st_mtime, mcfg.get("owner_match"))
    if key in _MIS_CACHE:
        return _MIS_CACHE[key], p.name
    try:
        import openpyxl
    except ImportError:
        return None, "openpyxl not installed"
    try:
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        ws = wb[mcfg.get("sheet", "Work items")]
        it = ws.iter_rows(values_only=True)
        hdr = list(next(it))
        ix = {h: i for i, h in enumerate(hdr) if h}
        owner_col, date_col = mcfg.get("owner_column", "Owners"), mcfg.get("date_column", "Created")
        if owner_col not in ix or date_col not in ix:
            return None, f"columns {owner_col}/{date_col} missing from {p.name}"
        want = mcfg.get("owner_match", "")
        fmt = mcfg.get("date_format", "%d/%m/%Y")
        redact = Redactor(cfg)
        rows = []
        for r in it:
            owners = str(r[ix[owner_col]] or "")
            if want and want not in owners:
                continue
            try:
                d = datetime.strptime(str(r[ix[date_col]]).strip(), fmt).date()
            except (ValueError, TypeError):
                continue

            def cell(name, limit=400):
                i = ix.get(name)
                v = r[i] if i is not None else None
                return redact(str(v).strip())[:limit] if v not in (None, "") else ""

            rows.append({"date": d.isoformat(), "title": cell("Title", 200),
                         "description": cell("Description"), "status": cell("Status", 40),
                         "priority": cell("Priority", 40), "type": cell("Type", 40),
                         "project": cell("Project", 80), "owners": redact(owners)[:120]})
        _MIS_CACHE[key] = rows
        return rows, p.name
    except Exception as exc:
        return None, f"{p.name} unreadable ({exc!r})"


def collect_mis_board(cfg: dict, day: date_cls) -> dict:
    """Tasks the user logged on the team board, dated by the board's own column.

    Hours Spent is deliberately NOT read: the column was added late and defaults
    to 1h, so it says nothing about effort.
    """
    mcfg = cfg.get("mis_board") or {}
    if not mcfg.get("enabled", False):
        return {"available": False, "reason": "disabled"}
    rows, origin = _mis_rows_for(cfg)
    if rows is None:
        return {"available": False, "reason": origin}
    mine = [r for r in rows if r["date"] == day.isoformat()]
    return {"available": True, "source": f"{origin} ({mcfg.get('date_column', 'Created')} date)",
            "count": len(mine), "tasks": mine,
            "note": "Hours Spent is not read: the column defaults to 1h and is unreliable"}


def build_digest(cfg: dict, day: date_cls) -> dict:
    start, end = window_bounds(cfg, day)
    attended = collect_meetings_attended(cfg, day)
    calendar = collect_calendar(cfg, day)
    return {
        "date": day.isoformat(),
        "weekday": day.strftime("%A"),
        "window": {"start": start.isoformat(), "end": end.isoformat()},
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "sources": {
            "activity": collect_activity(cfg, day),
            "git": collect_git(cfg, day),
            "shell": collect_shell(cfg, day),
            "calendar": calendar,
            "meetings_attended": attended,
            "unjoined_meetings": collect_unjoined_meetings(cfg, day, attended, calendar),
            "mis_board": collect_mis_board(cfg, day),
            "claude_code": collect_claude_code(cfg, day),
            "cloud": collect_cloud(cfg, day),
            "claude_export": collect_claude_export(cfg, day),
        },
    }


def digest_has_content(d: dict) -> bool:
    s = d["sources"]
    return any([
        s["activity"].get("apps"),
        s["git"].get("total_commits"),
        s["shell"].get("count"),
        s["calendar"].get("events"),
        s.get("mis_board", {}).get("count"),
        s.get("unjoined_meetings", {}).get("count"),
        s.get("meetings_attended", {}).get("count"),
        s["claude_code"].get("session_count"),
        s["cloud"].get("count"),
        s["claude_export"].get("count"),
    ])


# --------------------------------------------------------------------------- #
# markdown rendering
# --------------------------------------------------------------------------- #

SUMMARY_PROMPT = """You are writing a daily work log from machine-collected activity data.

Produce markdown in EXACTLY this shape:

**Work update — {date_h}**

**<Area heading>**
- <One-sentence past-tense bullet.>

Rules:
- One sentence per item, past tense, starting with the verb: "Fixed…", "Implemented…", "Reviewed…", "Investigated…".
- Group under short bold area headings (project or area names, e.g. **Backend**, **Infra**, **Meetings**). If everything is one area, skip headings and use a flat list.
- Include a number or named artefact as evidence when the data records one (commit subject, file name, doc name). Keep it inside the sentence.
- 3-10 bullets. Neutral and factual — audience is the user's manager. No self-praise, no "successfully", no process narration.
- Base every bullet on evidence in the JSON. Commit subjects, meeting names (meetings_attended and unjoined_meetings) and Claude Code prompts are the strongest signals of intent; app/window time and shell commands are supporting context only.
- When commits exist, describe what actually changed using the commit subject, body and changed-file stats. If more than one repo has commits, make clear per repo what was done (separate bullets or separate area headings per repo).
- Do NOT invent work. Do NOT list raw app names or timings as bullets. In the summary bullets
  keep file names out; inside the Evidence block naming real files, branches and tests is wanted,
  but use the short repo-relative form the JSON gives, never an absolute path.
- If a meeting happened, one bullet for meetings is enough unless they were clearly distinct pieces of work.
- Only write that a meeting was attended when it appears in `meetings_attended`. A meeting in
  `unjoined_meetings` was booked and ran but was NOT joined on this machine — word it as
  "did not join" or "booked but not joined", never as attended. Do not present
  `unjoined_meetings` as the day's calendar; it is explicitly incomplete.
- If the JSON has an `after_hours_included` key, the items it counts were done after the
  workday window closed. Fold them into the bullets, the Evidence and the MIS rows like any
  other work. Do NOT write a separate after-hours section — one is appended automatically.
- If the JSON has a `reconstructed` key, this day was rebuilt long after the fact. Its
  `missing_sources` list says what was never captured: there is no app/window or shell data
  and no record of which calls were actually joined. Write no `Focus time` section. Booked
  meetings from `calendar` may still be listed, but say only that they were booked — there is
  no attendance evidence for that day at all. Do not read a missing source as evidence that
  nothing happened, and infer Time Spent from session durations, commit volume and logged
  board tasks alone.
{next_up}
After the bullets, add this section verbatim, then the raw evidence block that follows the JSON marker:

---
**Evidence**

Then under Evidence, write compact sub-sections only for sources that have data:
- `Commits` — if one repo: bullet per commit as `subject (sha) — N files, +X/-Y`. If several
  repos: a bold repo name line, then its commit bullets. When a commit has a body, add one
  indented sub-bullet with a one-line gist of it. Mention branches touched per repo if recorded.
- `Meetings` — up to three sources, each answering a different question. `calendar` is what
  was booked, read from the published Outlook feed: list as `HH:MM–HH:MM — subject (booked)`,
  and mark an `out_of_office` entry as such. `meetings_attended` is the record of calls
  actually joined, whose `duration` is the full time the call was open and therefore the time
  really spent in it: list as `HH:MM–HH:MM — title (call joined, duration)`. When a booked
  meeting and a joined call clearly match, keep one bullet reading `booked, joined N`.
  `unjoined_meetings` is a weak fallback used only when the calendar has nothing for the day.
  A booked meeting with no matching joined call was NOT necessarily attended — say nothing
  about attending it. Never merge these lists or infer one from another.
- `MIS board tasks logged` — from `mis_board`, the tasks the user recorded on the team board
  that day, with status and priority. These are the user's own words for the work and are the
  strongest evidence of intent when commits are thin. List one bullet per task as
  `title (status)`. Hours are deliberately absent from this source; never invent them from it.
- `Claude Code` — this is usually the richest record of the day, so make it detailed.
  One bold sub-heading per session: `**<name> — <duration>, N turns**`. `<name>` MUST be the
  session's `title` field whenever it is non-empty — it is the label the work was given and it
  distinguishes sessions sharing a project; fall back to the project name only when `title` is
  empty. Append `also_titled` entries in brackets, and mark `subagent: true` as `(subagent)`.
  Skip a session with no `prompts`, no `did` and no `files_changed`: it records nothing worth a
  heading. Under each, write:
  - one sentence saying what the session set out to do, taken from `prompts` and `title`;
  - `Worked on:` a compact list of the concrete tasks, built from `did` (the intents of the
    commands actually run) and `files_changed`. Name real files and branches. Group related
    steps into one entry rather than transcribing every command;
  - `Tested:` only when `tests_touched` is non-empty or `did` shows tests being run — say
    what was verified;
  - `Status:` one of `completed`, `in progress`, or `parked`, decided from the evidence:
    a session whose `did` ends in a commit or a passing test run is completed; one whose
    `left_off` describes work still to do is in progress or parked. Quote or paraphrase
    `left_off` in a few words when it says where things stand.
  Sessions in the same project on the same day are separate sessions — keep them separate.
  Do not invent progress that the evidence does not show.
- `Files touched (OneDrive/SharePoint)` — up to 15 bullets of file names
- `Focus time` — one line listing top apps with durations

After Evidence, add a final section for the status tracker:

---
**{section}**

| Agent | Action Item | Owner(s) | Priority | Status | Time Spent | Description |
|---|---|---|---|---|---|---|---|

One data row per distinct piece of work from the bullets above — merge related bullets
into one row; 1-5 rows total. Do NOT wrap the table in a code fence. Column rules:
- Agent: which tracker project the work belongs to. This is a fixed dropdown — the value must
  be one of the strings below, copied EXACTLY, character for character. The list may contain
  what look like misspellings or odd spacing; they are NOT typos to correct. A "corrected"
  string does not match the dropdown and the card lands nowhere.
{agent_list}
  Work out the project from the evidence: the repository, the git branch, the files changed
  and the session titles all name it. If none of the entries above fits, do NOT force an entry
  and do NOT write a placeholder: write the name of the repository or area the work was
  actually in, as the evidence records it. That keeps the row filed against something real and
  makes it obvious the cell needs setting by hand. Separate multiple projects with ` / `.
  Never leave the cell blank.
{agent_hints}
- Action Item: short plain title, 4-10 words, states the problem or task, not the solution.
- Owner(s): `{owner}`; separate multiple owners with ` / `.
- Priority: High for blockers, production bugs or named feedback; otherwise Medium.
- Status: Done if the evidence shows the work completed/verified; In Progress if it was
  worked on but not clearly finished; To Do only for clearly recorded planned work.
- Time Spent: engineering effort in HOURS inferred from session durations and focus
  time — `2h`, `4h`, `0.5h`, never days; prefix `~` since it is inferred. For meeting
  time use the `duration` from `meetings_attended`, never the app's own foreground
  time, which understates a call spent working in another window.
- Description: ONE line, no line breaks, no literal pipe characters; compress
  problem, action taken and evidence into a single sentence joined by semicolons or an
  em-dash; include concrete numbers when the data records them; do not restate the title.

Activity JSON:
```json
{payload}
```

Output only the markdown. No preamble, no code fence around the whole thing."""

# The `claude -p` call this script makes is itself recorded as a Claude Code
# session, so without this the reporter reports its own summarizing as work.
_SELF_PROMPT_HEADS = (
    SUMMARY_PROMPT.split("\n", 1)[0],
    "Merge these daily MIS dashboard rows",
)


def is_self_prompt(txt: str) -> bool:
    head = " ".join((txt or "").split())[:120]
    return any(head.startswith(h[:60]) for h in _SELF_PROMPT_HEADS)


def _git_lines(g: dict) -> list[str]:
    if not g.get("total_commits"):
        return []
    L = ["**Commits**"]
    multi = len(g["repos_with_activity"]) > 1
    for repo in g["repos_with_activity"]:
        if not repo["commits"] and not repo["branches_touched"]:
            continue
        if multi:
            L.append(f"- **{repo['repo']}**")
        pad = "  " if multi else ""
        for c in repo["commits"]:
            stat = ""
            if c.get("files_changed"):
                stat = (f" — {c['files_changed']} file"
                        f"{'s' if c['files_changed'] != 1 else ''}, "
                        f"+{c['insertions']}/-{c['deletions']}")
            L.append(f"{pad}- {c['subject']} ({c['sha']}){stat}")
            if c.get("body"):
                gist = c["body"].splitlines()[0][:160]
                L.append(f"{pad}  - {gist}")
        if repo["branches_touched"]:
            L.append(f"{pad}- branches touched — "
                     f"{', '.join(repo['branches_touched'])}")
    L.append("")
    return L


def render_fallback(cfg: dict, d: dict) -> str:
    s = d["sources"]
    day_h = datetime.fromisoformat(d["date"]).strftime("%-d %b %Y")
    L = [f"**Work update — {day_h}**", ""]

    if not digest_has_content(d):
        L.append(f"No recorded activity found for {day_h}. "
                 "Check `worklog doctor` — the tracker may not have been running.")
        return "\n".join(L)

    L.append("_Generated without summarization (Claude CLI unavailable) — raw evidence below._")
    L.append("")
    L.append("---")
    L.append("**Evidence**")
    L.append("")

    L.extend(_git_lines(s["git"]))

    cal = s["calendar"]
    att = s.get("meetings_attended", {})
    sch = s.get("unjoined_meetings", {})
    if cal.get("events") or att.get("count") or sch.get("count"):
        L.append("**Meetings**")
        for e in cal.get("events") or []:
            if e.get("start", "")[11:16]:
                oof = ", out of office" if e.get("out_of_office") else ""
                L.append(f"- {e['start'][11:16]}–{e['end'][11:16]} — "
                         f"{e.get('subject', '')} (booked{oof})")
            else:
                L.append(f"- {e.get('subject') or e.get('raw', '')}")
        for m in sch.get("meetings") or []:
            L.append(f"- {m['start'][11:16]}–{m['end'][11:16]} — {m['title']} "
                     f"(booked, not joined)")
        for m in att.get("meetings") or []:
            L.append(f"- {m['start'][11:16]}–{m['end'][11:16]} — {m['title']} "
                     f"(call joined, {m['duration']})")
        L.append("")

    mis = s.get("mis_board") or {}
    if mis.get("count"):
        L.append("**MIS board tasks logged**")
        for t in mis["tasks"]:
            bits = " · ".join(x for x in (t.get("status"), t.get("priority"),
                                          t.get("project")) if x)
            L.append(f"- {t['title']}" + (f" ({bits})" if bits else ""))
        L.append("")

    cc = s["claude_code"]
    if cc.get("session_count"):
        L.append("**Claude Code**")
        for sess in cc["sessions"]:
            proj = Path(sess["project"]).name or sess["project"]
            head = sess.get("title") or proj
            extra = f" · {proj}" if sess.get("title") else ""
            if sess.get("subagent"):
                extra += " · subagent"
            L.append(f"- **{head}** — {sess['duration']}, "
                     f"{sess['assistant_turns']} turns{extra}")
            if sess.get("branches"):
                L.append(f"  - branch: {', '.join(sess['branches'])}")
            for p in sess["prompts"][:5]:
                L.append(f"  - asked: {p}")
            for d in (sess.get("did") or [])[:10]:
                L.append(f"  - did: {d}")
            if sess.get("files_changed"):
                more = sess.get("files_changed_total", 0) - len(sess["files_changed"])
                tail = f" …and {more} more" if more > 0 else ""
                L.append(f"  - changed: {', '.join(sess['files_changed'])}{tail}")
            if sess.get("tests_touched"):
                L.append(f"  - tests: {', '.join(sess['tests_touched'])}")
            if sess.get("skills_used"):
                L.append(f"  - skills: {', '.join(sess['skills_used'])}")
            if sess.get("left_off"):
                L.append(f"  - left off: {sess['left_off']}")
        L.append("")

    cl = s["cloud"]
    if cl.get("count"):
        L.append("**Files touched (OneDrive/SharePoint)**")
        for f in cl["files"][:15]:
            L.append(f"- {f['file']}")
        if cl["count"] > 15:
            L.append(f"- …and {cl['count'] - 15} more")
        L.append("")

    ex = s["claude_export"]
    if ex.get("count"):
        L.append("**Claude.ai conversations**")
        for c in ex["conversations"]:
            L.append(f"- {c['name']}")
        L.append("")

    act = s["activity"]
    if act.get("apps"):
        top = ", ".join(f"{a['app']} {a['duration']}" for a in act["apps"][:6])
        L.append("**Focus time**")
        L.append(f"- {top} (active {act['active_duration']}, idle {act['idle_duration']})")
        L.append("")

    sh = s["shell"]
    if sh.get("count"):
        L.append("**Shell**")
        for c in sh["commands"][:20]:
            L.append(f"- `{c['cmd']}`")
        L.append("")

    L.append("---")
    L.append(SOR_MARKER)
    L.append("")
    L.append("| Agent | Action Item | Owner(s) | Priority | Status | Time Spent | Description |")
    L.append("|---|---|---|---|---|---|---|---|")
    L.append("")
    L.append("_Claude CLI was unavailable — fill rows from the evidence above._")

    return "\n".join(L).rstrip() + "\n"


def summarize(cfg: dict, d: dict) -> tuple[str, str]:
    """Returns (markdown, method)."""
    scfg = cfg["summarizer"]
    if scfg["mode"] != "claude_cli":
        return render_fallback(cfg, d), "fallback"

    binary = shutil.which(scfg["claude_binary"]) or scfg["claude_binary"]
    if not shutil.which(binary):
        log_line(cfg, "summarize: claude CLI not on PATH, using fallback")
        return render_fallback(cfg, d), "fallback"

    day_h = datetime.fromisoformat(d["date"]).strftime("%-d %b %Y")
    next_up = ("- Add a final **Next up** section only for work clearly recorded as "
               "an agreed next step.\n") if scfg["include_next_up"] else ""
    payload = json.dumps(d, ensure_ascii=False, indent=1)
    if len(payload) > 180_000:
        log_line(cfg, f"summarize: digest is {len(payload)} chars, "
                      f"dropping {len(payload) - 180_000} — evidence will be incomplete")
    tcfg = cfg.get("tracker") or {}
    agents = tcfg.get("agents") or []
    hints = (tcfg.get("agent_hints") or "").strip()
    agent_hints = ("  " + hints) if hints else ""
    owner = tcfg.get("default_owner") or "me"
    agent_list = ("\n".join(f"    {a}" for a in agents) if agents else
                  "    (no project list configured — see tracker.agents in the config)")
    prompt = SUMMARY_PROMPT.format(
        date_h=day_h, next_up=next_up, agent_list=agent_list,
        agent_hints=agent_hints, owner=owner,
        section=tcfg.get("section_title") or DEFAULT_SECTION,
        payload=payload[:180_000])

    cmd = [binary, "-p", prompt]
    if scfg.get("model"):
        cmd += ["--model", scfg["model"]]

    # A usage limit or a dropped connection is transient, and the cost of giving up
    # is a whole day rendered without summarization. Retry before falling back.
    attempts = max(int(scfg.get("attempts", 3)), 1)
    delay = float(scfg.get("retry_delay_seconds", 30))
    rc, out, err = 0, "", ""
    for n in range(1, attempts + 1):
        rc, out, err = run(cmd, timeout=scfg["timeout_seconds"])
        if rc == 0 and out.strip():
            break
        # claude reports some errors on stdout, so log both or the reason is invisible
        detail = (err.strip() or out.strip() or "no output")[:300]
        log_line(cfg, f"summarize: claude CLI failed rc={rc} attempt {n}/{attempts} "
                      f"— {detail}")
        if n < attempts:
            time.sleep(delay)
    if rc != 0 or not out.strip():
        return render_fallback(cfg, d), "fallback"

    md = out.strip()
    if md.startswith("```"):
        md = re.sub(r"^```[a-z]*\n", "", md)
        md = re.sub(r"\n```$", "", md)
    return md.strip() + "\n", "claude_cli"


# --------------------------------------------------------------------------- #
# after-hours addendum: work done after the previous workday's report ran
# --------------------------------------------------------------------------- #

ADDENDUM_MARKER = "**After-hours addendum**"


def build_addendum(cfg: dict, prev_day: date_cls, today: date_cls) -> dict:
    """Everything between prev workday's window end and today's window start."""
    _, prev_end = window_bounds(cfg, prev_day)
    today_start, _ = window_bounds(cfg, today)
    bounds = (prev_end, today_start)
    return {
        "date": prev_day.isoformat(),
        "kind": "addendum",
        "window": {"start": prev_end.isoformat(), "end": today_start.isoformat()},
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "sources": {
            "git": collect_git(cfg, prev_day, bounds),
            "shell": collect_shell(cfg, prev_day, bounds),
            "claude_code": collect_claude_code(cfg, prev_day, bounds),
            "cloud": collect_cloud(cfg, prev_day, bounds),
        },
    }


def addendum_headline(ad: dict) -> str:
    s = ad["sources"]
    bits = []
    for key, noun in (("git", "commit"), ("claude_code", "CC session"),
                      ("shell", "shell command"), ("cloud", "file")):
        n = (s[key].get("total_commits") if key == "git"
             else s[key].get("session_count") if key == "claude_code"
             else s[key].get("count")) or 0
        if n:
            bits.append(f"{n} {noun}{'s' if n != 1 else ''}")
    return " · ".join(bits)


def render_addendum(ad: dict) -> str:
    s = ad["sources"]
    start = datetime.fromisoformat(ad["window"]["start"])
    added = datetime.fromisoformat(ad["generated_at"]).strftime("%-d %b")
    L = ["", "---",
         f"{ADDENDUM_MARKER} _(activity after {start.strftime('%H:%M on %-d %b')}, "
         f"added {added})_", ""]

    L.extend(_git_lines(s["git"]))

    cc = s["claude_code"]
    if cc.get("session_count"):
        L.append("**Claude Code**")
        for sess in cc["sessions"]:
            proj = Path(sess["project"]).name or sess["project"]
            L.append(f"- {proj} — {sess['duration']}, {sess['assistant_turns']} turns")
            for p in sess["prompts"][:3]:
                L.append(f"  - {p}")
        L.append("")

    cl = s["cloud"]
    if cl.get("count"):
        L.append("**Files touched (OneDrive/SharePoint)**")
        for f in cl["files"][:15]:
            L.append(f"- {f['file']}")
        L.append("")

    sh = s["shell"]
    if sh.get("count"):
        L.append("**Shell**")
        for c in sh["commands"][:20]:
            L.append(f"- `{c['cmd']}`")
        L.append("")

    return "\n".join(L).rstrip() + "\n"


def merge_addendum(d: dict, ad: dict) -> dict:
    """Fold after-hours items into a day's digest so the MIS rows cover them.

    The addendum block appended to the file records *when* the work happened; this
    makes the summarized body and the dashboard table include *that* it happened.
    """
    m = json.loads(json.dumps(d))
    src, asrc, counts = m["sources"], ad["sources"], {}

    g, ag = src.get("git") or {}, asrc.get("git") or {}
    if ag.get("total_commits"):
        by_repo = {r["repo"]: r for r in g.get("repos_with_activity") or []}
        for ar in ag.get("repos_with_activity") or []:
            r = by_repo.get(ar["repo"])
            if r is None:
                g.setdefault("repos_with_activity", []).append(ar)
            else:
                r["commits"] = (r.get("commits") or []) + (ar.get("commits") or [])
                r["branches_touched"] = sorted(set(r.get("branches_touched") or [])
                                               | set(ar.get("branches_touched") or []))
        g["available"] = True
        g["total_commits"] = (g.get("total_commits") or 0) + ag["total_commits"]
        counts["commits"] = ag["total_commits"]

    cc, acc = src.get("claude_code") or {}, asrc.get("claude_code") or {}
    if acc.get("session_count"):
        cc["available"] = True
        cc["sessions"] = (cc.get("sessions") or []) + (acc.get("sessions") or [])
        cc["session_count"] = (cc.get("session_count") or 0) + acc["session_count"]
        counts["claude_code_sessions"] = acc["session_count"]

    for key, items in (("shell", "commands"), ("cloud", "files")):
        s, a = src.get(key) or {}, asrc.get(key) or {}
        if a.get("count"):
            s["available"] = True
            s[items] = (s.get(items) or []) + (a.get(items) or [])
            s["count"] = (s.get("count") or 0) + a["count"]
            counts[key] = a["count"]

    if counts:
        m["after_hours_included"] = {"window": ad["window"], **counts}
    return m


def stored_addendum(cfg: dict, day: date_cls) -> dict | None:
    p = expand(cfg["state_dir"]) / "raw" / day.isoformat() / "addendum.json"
    if not p.is_file():
        return None
    try:
        return json.loads(p.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError):
        return None


def apply_addendum(cfg: dict, today: date_cls) -> str | None:
    """Append late work to the previous workday's log. Returns headline or None."""
    prev = prev_workday(cfg, today)
    if prev is None:
        return None
    out_path = expand(cfg["output_dir"]) / f"{prev.isoformat()}.md"
    if out_path.is_file() and ADDENDUM_MARKER in out_path.read_text(encoding="utf-8"):
        return None
    ad = build_addendum(cfg, prev, today)
    headline = addendum_headline(ad)
    if not headline:
        return None

    md = render_addendum(ad)
    if out_path.is_file():
        with open(out_path, "a", encoding="utf-8") as fh:
            fh.write(md)
    else:
        day_h = datetime.fromisoformat(ad["date"]).strftime("%-d %b %Y")
        out_path.parent.mkdir(parents=True, exist_ok=True)
        out_path.write_text(
            f"<!-- worklog addendum-only | generated {ad['generated_at']} -->\n\n"
            f"**Work update — {day_h}**\n\n"
            f"_No log was generated for this day; after-hours activity below._\n" + md,
            encoding="utf-8")

    raw_dir = expand(cfg["state_dir"]) / "raw" / prev.isoformat()
    raw_dir.mkdir(parents=True, exist_ok=True)
    (raw_dir / "addendum.json").write_text(
        json.dumps(ad, ensure_ascii=False, indent=2), encoding="utf-8")
    log_line(cfg, f"addendum: appended to {out_path} — {headline}")
    return f"{headline} → {prev.strftime('%-d %b')}"


# --------------------------------------------------------------------------- #
# backfill: redo days whose log is missing or was written without the LLM
# --------------------------------------------------------------------------- #

REPAIR_DEFAULTS = {"enabled": True, "lookback_days": 7, "max_days_per_run": 3}
DEFAULT_SECTION = "Status Tracker"
SOR_MARKER = "**Status Tracker**"
# The section was called this before the tracker was renamed. Both are accepted so the
# reports already written stay healthy instead of all looking degraded at once and
# dragging the whole backlog through a needless re-summarization.
MIS_MARKER = "**MIS Board Tasks**"
TRACKER_MARKERS = (SOR_MARKER, MIS_MARKER, "**SOR Status Tracker**")
AFTER_HOURS_TAG = "+after-hours"
RECONSTRUCTED_NOTE = (
    "_Reconstructed after the fact from git history, Claude Code sessions, OneDrive "
    "file times, the published Outlook calendar and the MIS board export. The activity "
    "tracker did not run on this day, so app/window time, shell history and which calls "
    "were actually joined are unavailable._\n\n")
_HEADER_RE = re.compile(r"<!--\s*worklog\s+([a-z_-]+)")


def repair_cfg(cfg: dict) -> dict:
    return {**REPAIR_DEFAULTS, **(cfg.get("repair") or {})}


def report_health(path: Path) -> tuple[str, str]:
    """Returns ('ok'|'missing'|'degraded', reason)."""
    if not path.is_file():
        return "missing", "no log file"
    text = path.read_text(encoding="utf-8")
    header = text.split("\n", 1)[0]
    m = _HEADER_RE.search(header)
    method = m.group(1) if m else "unknown"
    if method != "claude_cli":
        return "degraded", f"written via {method}"
    if not any(m in text for m in TRACKER_MARKERS):
        return "degraded", "no SOR Status Tracker section"
    if ADDENDUM_MARKER in text and AFTER_HOURS_TAG not in header:
        return "degraded", "after-hours work missing from the MIS rows"
    return "ok", method


def _addendum_tail(text: str) -> str:
    """The trailing after-hours block of an existing log, or ''."""
    i = text.find(ADDENDUM_MARKER)
    if i < 0:
        return ""
    sep = text.rfind("\n---\n", 0, i)
    return text[sep:] if sep >= 0 else "\n---\n" + text[i:]


# Only these two cannot be re-derived for a past day: shell history rotates and had no
# timestamps before the EXTENDED_HISTORY change, and cloud detection is by file mtime,
# which now reports today. Everything else — the activity samples, git, the ICS feed, the
# board export, the Claude Code transcripts — is still on disk, so a repair re-collects it
# and picks up whatever the collectors learned since the day ran.
LOSSY_SOURCES = ("shell", "cloud")


def load_digest(cfg: dict, day: date_cls) -> dict | None:
    """A digest for a past day: freshly collected, with the lossy parts preserved."""
    p = expand(cfg["state_dir"]) / "raw" / day.isoformat() / "digest.json"
    stored = None
    if p.is_file():
        try:
            stored = json.loads(p.read_text(encoding="utf-8"))
        except (json.JSONDecodeError, OSError) as exc:
            log_line(cfg, f"repair: unreadable {p} ({exc!r}), rebuilding")

    if stored:
        try:
            fresh = build_digest(cfg, day)
        except Exception as exc:            # a failed refresh must not lose the day
            log_line(cfg, f"repair: refresh failed for {day} ({exc!r}), using stored")
            return stored
        for key in LOSSY_SOURCES:
            if key in (stored.get("sources") or {}):
                fresh["sources"][key] = stored["sources"][key]
        for key in ("reconstructed", "after_hours_included"):
            if key in stored:
                fresh[key] = stored[key]
        return fresh if digest_has_content(fresh) else stored

    d = build_digest(cfg, day)
    if not digest_has_content(d):
        return None
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(json.dumps(d, ensure_ascii=False, indent=2), encoding="utf-8")
    return d


def repair_day(cfg: dict, day: date_cls) -> str | None:
    """Re-summarize one day in place, keeping its addendum. Returns a note."""
    d = load_digest(cfg, day)
    if d is None:
        log_line(cfg, f"repair: {day} has no recoverable activity, skipped")
        return None

    tag = ""
    ad = stored_addendum(cfg, day)
    if ad:
        merged = merge_addendum(d, ad)
        if "after_hours_included" in merged:
            d, tag = merged, f" {AFTER_HOURS_TAG}"

    md, method = summarize(cfg, d)
    if method != "claude_cli":
        log_line(cfg, f"repair: {day} still unsummarized (claude CLI down), kept as is")
        return None

    out_path = expand(cfg["output_dir"]) / f"{day.isoformat()}.md"
    tail = _addendum_tail(out_path.read_text(encoding="utf-8")) if out_path.is_file() else ""
    # a reconstructed day was never gated to office hours and has no tracker data;
    # repairing it must not restate the live window or lose that provenance
    recon = d.get("reconstructed") or {}
    window = "full day" if recon else f"{cfg['work_hours']['start']}–{cfg['work_hours']['end']}"
    header = (f"<!-- worklog {method} repaired{' backfilled' if recon else ''}{tag} | "
              f"generated {datetime.now().isoformat(timespec='seconds')} | "
              f"window {window} -->\n\n")
    if recon:
        md = RECONSTRUCTED_NOTE + md if RECONSTRUCTED_NOTE.strip() not in md else md
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(header + md.rstrip("\n") + "\n" + tail, encoding="utf-8")
    # the stored digest is deliberately NOT rewritten here: `d` may be merged with
    # the addendum, and saving that would double-count on the next repair
    log_line(cfg, f"repair: rewrote {out_path} via {method}{tag}")
    return day.strftime("%-d %b")


def first_tracked_day(cfg: dict) -> date_cls | None:
    """Earliest day the tracker has raw data for — nothing before it is fixable."""
    raw = expand(cfg["state_dir"]) / "raw"
    days = sorted(p.name for p in raw.glob("20[0-9][0-9]-[0-9][0-9]-[0-9][0-9]")
                  if p.is_dir())
    return datetime.strptime(days[0], "%Y-%m-%d").date() if days else None


def scan_backlog(cfg: dict, today: date_cls) -> list[tuple]:
    """Workdays in the lookback window whose log is missing or degraded."""
    rc = repair_cfg(cfg)
    floor = first_tracked_day(cfg)
    out_dir = expand(cfg["output_dir"])
    found = []
    for i in range(1, int(rc["lookback_days"]) + 1):
        day = today - timedelta(days=i)
        if floor and day < floor:
            continue
        path = out_dir / f"{day.isoformat()}.md"
        # weekends are not scanned for missing logs, but a weekend day that already
        # has one — from weekend work or a backfill — must still be kept correct
        if day.isoweekday() not in cfg["work_days"] and not path.is_file():
            continue
        status, reason = report_health(path)
        if status != "ok":
            found.append((day, status, reason))
    return sorted(found)


def run_repairs(cfg: dict, today: date_cls) -> str | None:
    """Fix earlier days before today's report is announced. Returns a note."""
    rc = repair_cfg(cfg)
    if not rc["enabled"]:
        return None
    backlog = scan_backlog(cfg, today)
    if not backlog:
        return None
    cap = int(rc["max_days_per_run"])
    for day, status, reason in backlog[cap:]:
        log_line(cfg, f"repair: deferred {day} ({status}: {reason}) — "
                      f"max_days_per_run={cap}")
    fixed = []
    for day, status, reason in backlog[:cap]:
        log_line(cfg, f"repair: {day} {status} ({reason}), re-summarizing")
        try:
            note = repair_day(cfg, day)
        except Exception as exc:  # one bad day must not stop the others
            log_line(cfg, f"repair: {day} failed {exc!r}")
            continue
        if note:
            fixed.append(note)
    return ", ".join(fixed) or None


# --------------------------------------------------------------------------- #
# tracker health: a silent sampler produces a hollow log that repair cannot fix
# --------------------------------------------------------------------------- #

HEALTH_DEFAULTS = {"enabled": True, "sample_interval_seconds": 60,
                   "min_sample_ratio": 0.5}


def sampler_health(cfg: dict, day: date_cls, now: datetime | None = None) -> dict | None:
    """Compare samples taken against samples expected. None when not applicable."""
    hcfg = {**HEALTH_DEFAULTS, **(cfg.get("health") or {})}
    if not hcfg["enabled"] or day.isoweekday() not in cfg["work_days"]:
        return None
    now = now or datetime.now()
    start, end = window_bounds(cfg, day)
    elapsed = ((min(now, end) if day == now.date() else end) - start).total_seconds() / 60.0
    if elapsed < 30:                      # too early to judge
        return None
    # paused minutes were never meant to be sampled, so charging them against the
    # expected count would raise a false "tracker is down" alarm every time
    paused = paused_minutes(cfg, day, start, min(now, end) if day == now.date() else end)
    elapsed = max(elapsed - paused, 0.0)
    interval = max(float(hcfg["sample_interval_seconds"]), 1.0) / 60.0
    expected = elapsed / interval
    path = expand(cfg["state_dir"]) / "raw" / day.isoformat() / "activity.jsonl"
    got = sum(1 for _ in open(path, errors="replace")) if path.is_file() else 0
    ratio = got / expected if expected else 1.0
    return {"samples": got, "expected": int(expected), "ratio": round(ratio, 2),
            "paused_minutes": round(paused, 1),
            "ok": ratio >= float(hcfg["min_sample_ratio"])}


# --------------------------------------------------------------------------- #
# weekly rollup: one MIS table for the week, to paste into the dashboard
# --------------------------------------------------------------------------- #

WEEKLY_PROMPT = """Merge these daily MIS dashboard rows into one table for the week.

Output ONLY a markdown table with exactly these columns and no other text:

| Agent | Action Item | Owner(s) | Priority | Status | Time Spent | Description |
|---|---|---|---|---|---|---|---|

Rules:
- Merge rows describing the same piece of work across days into ONE row. Sum their hours.
- Keep rows for distinct work separate. 3-12 rows total.
- Owner(s): `{owner}`.
- Agent: copy the value from the daily rows EXACTLY, character for character, including any
  apparent misspelling or odd spacing. Never correct one. Only merge rows that share the same Agent;
  the same-sounding work under two different Agents stays two rows. Some daily rows predate
  this column, and some also carry a leading card-number cell that no longer exists, so a
  daily row may have six, seven or eight cells. Normalize each one to the columns above:
  discard any leading card-number cell, and where the Agent is absent infer it from the Action
  Item and Description, falling back to the repository or area name if no entry fits.
- Priority: the highest seen among merged rows.
- Status: Done only if the latest day says Done; otherwise In Progress (or To Do if never started).
- Time Spent: summed engineering hours, prefixed `~`, e.g. `~6h`. Never days.
- Description: ONE line, no line breaks, no literal pipe characters; combine what the merged
  rows say, keeping concrete numbers; do not restate the Action Item.

Daily rows:
{rows}"""


def week_sat_to_fri(day: date_cls) -> tuple[date_cls, date_cls]:
    """The Saturday-to-Friday week containing `day`, as the user counts weeks."""
    friday = day + timedelta(days=(4 - day.weekday()) % 7)
    return friday - timedelta(days=6), friday


def write_mis_week(cfg: dict, day: date_cls) -> str | None:
    """Copy the MIS board export, keeping only rows worked on Sat-Fri this week.

    Row layout, columns and styling are the board's own: the workbook is copied and
    non-matching rows are deleted, so the result pastes back into the same tool.
    A row counts as worked on if it was created or updated inside the week,
    whatever its status.
    """
    mcfg = cfg.get("mis_board") or {}
    if not mcfg.get("weekly_export", True):
        return None
    src = expand(mcfg.get("path", "~/.worklog/mis-board.xlsx"))
    if not src.is_file():
        log_line(cfg, f"mis week: {src} not found, skipped")
        return None
    try:
        import openpyxl
    except ImportError:
        log_line(cfg, "mis week: openpyxl not installed, skipped")
        return None

    lo, hi = week_sat_to_fri(day)
    fmt = mcfg.get("date_format", "%d/%m/%Y")
    want = mcfg.get("owner_match", "")

    # the export is a manual snapshot; if it predates the week it cannot contain it
    stale = datetime.fromtimestamp(src.stat().st_mtime).date() < hi
    wb = openpyxl.load_workbook(src)
    ws = wb[mcfg.get("sheet", "Work items")]
    hdr = [c.value for c in ws[1]]
    ix = {h: i + 1 for i, h in enumerate(hdr) if h}
    owner_col = ix.get(mcfg.get("owner_column", "Owners"))
    if not owner_col:
        log_line(cfg, "mis week: owner column missing, skipped")
        return None

    def cell_date(row, name):
        col = ix.get(name)
        if not col:
            return None
        try:
            return datetime.strptime(str(ws.cell(row, col).value).strip(), fmt).date()
        except (ValueError, TypeError):
            return None

    keep, drop = 0, []
    for r in range(ws.max_row, 1, -1):
        owners = str(ws.cell(r, owner_col).value or "")
        created, updated = cell_date(r, "Created"), cell_date(r, "Updated")
        touched = any(d and lo <= d <= hi for d in (created, updated))
        if (not want or want in owners) and touched:
            keep += 1
        else:
            drop.append(r)
    for r in drop:                      # bottom-up, so indices stay valid
        ws.delete_rows(r)

    out_dir = expand(cfg["output_dir"])
    iso = hi.isocalendar()
    name = f"mis-week-{iso.year}-W{iso.week:02d}.xlsx"
    out_dir.mkdir(parents=True, exist_ok=True)
    wb.save(out_dir / name)
    log_line(cfg, f"mis week: wrote {name} — {keep} row(s) worked on "
                  f"{lo.isoformat()}..{hi.isoformat()}"
                  + (f" (WARNING: {src.name} last saved "
                     f"{datetime.fromtimestamp(src.stat().st_mtime).date()}, "
                     f"before the week ended — re-export for complete data)"
                     if stale else ""))
    return f"{name} ({keep} row{'s' if keep != 1 else ''})"


def week_days(cfg: dict, day: date_cls) -> list[date_cls]:
    """The workdays of `day`'s ISO week up to and including `day`."""
    monday = day - timedelta(days=day.isoweekday() - 1)
    return [d for i in range(7)
            if (d := monday + timedelta(days=i)) <= day
            and d.isoweekday() in cfg["work_days"]]


def _mis_rows(text: str) -> list[str]:
    """Data rows of a log's MIS table — the body rows, not header or separator."""
    i = min((j for j in (text.find(m) for m in TRACKER_MARKERS) if j >= 0), default=-1)
    if i < 0:
        return []
    rows = []
    for line in text[i:].splitlines():
        line = line.strip()
        if line == "---":            # start of the appended addendum block
            break
        if not line.startswith("|"):
            continue
        if set(line) <= set("|- ") or line.lower().startswith(("| s. no", "| agent")):
            continue
        rows.append(line)
    return rows


def is_last_workday_of_week(cfg: dict, day: date_cls) -> bool:
    return day.isoweekday() == max(cfg["work_days"])


def write_weekly(cfg: dict, day: date_cls) -> str | None:
    """Write YYYY-Www.md merging the week's MIS rows. Returns the file name."""
    out_dir = expand(cfg["output_dir"])
    days = week_days(cfg, day)
    per_day = []
    for d in days:
        p = out_dir / f"{d.isoformat()}.md"
        if not p.is_file():
            continue
        rows = _mis_rows(p.read_text(encoding="utf-8"))
        if rows:
            per_day.append((d, rows))
    if not per_day:
        log_line(cfg, f"weekly: no MIS rows in {days[0]}..{day}, skipped")
        return None

    blob = "\n\n".join(f"{d.strftime('%A %-d %b')}:\n" + "\n".join(r)
                       for d, r in per_day)
    scfg = cfg["summarizer"]
    binary = shutil.which(scfg["claude_binary"])
    table, method = None, "concatenated"
    if binary:
        rc, out, err = run([binary, "-p", WEEKLY_PROMPT.format(
                               rows=blob,
                               owner=(cfg.get("tracker") or {}).get("default_owner") or "me")],
                           timeout=scfg["timeout_seconds"])
        if rc == 0 and "|" in out:
            body = re.sub(r"^```[a-z]*\n|\n```$", "", out.strip())
            table, method = body.strip(), "claude_cli"
        else:
            log_line(cfg, f"weekly: claude CLI failed rc={rc} {err[:200]}")
    if table is None:
        table = "\n".join(
            ["| Agent | Action Item | Owner(s) | Priority | Status | Time Spent | Description |",
             "|---|---|---|---|---|---|---|---|"]
            + [r for _, rows in per_day for r in rows])

    iso = day.isocalendar()
    name = f"{iso.year}-W{iso.week:02d}.md"
    span = f"{per_day[0][0].strftime('%-d %b')}–{per_day[-1][0].strftime('%-d %b %Y')}"
    out_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / name).write_text(
        f"<!-- worklog weekly {method} | generated "
        f"{datetime.now().isoformat(timespec='seconds')} -->\n\n"
        f"**{(cfg.get('tracker') or {}).get('section_title') or DEFAULT_SECTION} — week {iso.week}, {span}**\n\n{table}\n\n"
        f"_Merged from {len(per_day)} daily log(s): "
        f"{', '.join(d.isoformat() for d, _ in per_day)}._\n",
        encoding="utf-8")
    log_line(cfg, f"weekly: wrote {name} via {method} from {len(per_day)} day(s)")
    return name


# --------------------------------------------------------------------------- #
# commands
# --------------------------------------------------------------------------- #

def report_lock(cfg: dict):
    """Hold an exclusive lock for the life of the process, or None if held already.

    Stops a manual `worklog report` and the 17:00 timer from summarizing the same
    day concurrently and racing each other's write.
    """
    path = expand(cfg["state_dir"]) / "report.lock"
    path.parent.mkdir(parents=True, exist_ok=True)
    fh = open(path, "w")
    try:
        fcntl.flock(fh, fcntl.LOCK_EX | fcntl.LOCK_NB)
    except OSError:
        fh.close()
        return None
    return fh          # deliberately never closed: the OS drops it on exit


def resolve_day(arg: str | None) -> date_cls:
    if not arg or arg == "today":
        return date_cls.today()
    if arg == "yesterday":
        return date_cls.today() - timedelta(days=1)
    try:
        return datetime.strptime(arg, "%Y-%m-%d").date()
    except ValueError:
        die(f"bad --date {arg!r}; use YYYY-MM-DD, 'today' or 'yesterday'")


def cmd_report(cfg: dict, args) -> int:
    day = resolve_day(args.date)
    lock = report_lock(cfg)
    if lock is None:
        log_line(cfg, "report: another run holds the lock, exiting")
        return 0
    d = build_digest(cfg, day)
    md, method = summarize(cfg, d)

    out_dir = expand(cfg["output_dir"])
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / f"{day.isoformat()}.md"
    header = (f"<!-- worklog {method} | generated {d['generated_at']} | "
              f"window {cfg['work_hours']['start']}–{cfg['work_hours']['end']} -->\n\n")
    out_path.write_text(header + md, encoding="utf-8")

    raw_dir = expand(cfg["state_dir"]) / "raw" / day.isoformat()
    raw_dir.mkdir(parents=True, exist_ok=True)
    (raw_dir / "digest.json").write_text(
        json.dumps(d, ensure_ascii=False, indent=2), encoding="utf-8")

    log_line(cfg, f"report: wrote {out_path} via {method} — {digest_headline(d)}")

    addendum_note = repair_note = weekly_note = None
    if day == date_cls.today():
        # addendum first: it stores addendum.json, which then makes the day look
        # degraded to report_health, so the repair pass folds it into the MIS rows
        try:
            addendum_note = apply_addendum(cfg, day)
        except Exception as exc:  # a broken catch-up must not kill the report
            log_line(cfg, f"addendum: failed {exc!r}")
        try:
            repair_note = run_repairs(cfg, day)
        except Exception as exc:  # a broken backfill must not kill the report
            log_line(cfg, f"repair: scan failed {exc!r}")

    if (cfg.get("weekly") or {}).get("enabled", True) and is_last_workday_of_week(cfg, day):
        try:
            weekly_note = write_weekly(cfg, day)
        except Exception as exc:  # a broken rollup must not kill the report
            log_line(cfg, f"weekly: failed {exc!r}")
        try:
            xl = write_mis_week(cfg, day)
            if xl:
                weekly_note = f"{weekly_note} + {xl}" if weekly_note else xl
        except Exception as exc:  # a broken export must not kill the report
            log_line(cfg, f"mis week: failed {exc!r}")

    health = sampler_health(cfg, day)
    if health and not health["ok"]:
        log_line(cfg, f"health: only {health['samples']} of ~{health['expected']} "
                      f"samples ({health['ratio']}) — tracker may be down")

    if not args.quiet:
        day_h = datetime.fromisoformat(d["date"]).strftime("%-d %b")
        if health and not health["ok"]:
            notify(cfg, "Work log: tracker may be down",
                   f"Only {health['samples']} of ~{health['expected']} samples today "
                   f"— app and meeting time is unreliable. Run `worklog doctor`.",
                   subtitle=day_h, open_path=str(out_path))
        elif digest_has_content(d):
            body = digest_headline(d)
            if addendum_note:
                body += f" · catch-up: {addendum_note}"
            if repair_note:
                body += f" · repaired: {repair_note}"
            if weekly_note:
                body += f" · weekly: {weekly_note}"
            notify(cfg, "Work log ready", body,
                   subtitle=f"{day_h} · {out_path.name}", open_path=str(out_path))
        else:
            notify(cfg, "Work log: nothing captured",
                   "Run `worklog doctor` — the tracker may not be running.",
                   subtitle=day_h, open_path=str(out_path))

    print(str(out_path))
    return 0


def cmd_repair(cfg: dict, args) -> int:
    out_dir = expand(cfg["output_dir"])
    if args.date:
        day = resolve_day(args.date)
        rows = [(day, *report_health(out_dir / f"{day.isoformat()}.md"))]
    else:
        rows = scan_backlog(cfg, date_cls.today())
        if not rows:
            print("nothing to repair")
            return 0
    for day, status, reason in rows:
        print(f"{day} — {status} ({reason})")
    if args.dry_run:
        return 0
    fixed, failed = [], []
    for day, status, reason in rows:
        if status == "ok" and not args.date:
            continue
        note = repair_day(cfg, day)
        print(f"{day} — {'repaired' if note else 'unchanged'}")
        (fixed if note else failed).append(day.strftime("%-d %b"))

    if not getattr(args, "quiet", False) and (fixed or failed):
        if fixed:
            body = f"Summarized {', '.join(fixed)}"
            if failed:
                body += f" · still waiting on {', '.join(failed)}"
            notify(cfg, "Work log caught up", body,
                   subtitle=f"{len(fixed)} day(s) repaired")
        else:
            notify(cfg, "Work log still needs the Claude CLI",
                   f"{', '.join(failed)} could not be summarized — will retry.",
                   subtitle="no connection?")
    return 0


def cmd_weekly(cfg: dict, args) -> int:
    name = write_weekly(cfg, resolve_day(args.date))
    print(name or "no MIS rows found for that week")
    return 0


def cmd_digest(cfg: dict, args) -> int:
    print(json.dumps(build_digest(cfg, resolve_day(args.date)),
                     ensure_ascii=False, indent=2))
    return 0


def cmd_doctor(cfg: dict, args) -> int:
    day = resolve_day(args.date)
    print(f"worklog doctor — {day.isoformat()}\n")
    print(f"output_dir : {expand(cfg['output_dir'])} "
          f"({'ok' if expand(cfg['output_dir']).parent.is_dir() else 'PARENT MISSING'})")
    print(f"state_dir  : {expand(cfg['state_dir'])}")
    print(f"claude CLI : {shutil.which(cfg['summarizer']['claude_binary']) or 'NOT FOUND'}")
    print(f"icalBuddy  : {shutil.which('icalBuddy') or 'not installed (optional)'}")
    print(f"notify via : {'terminal-notifier' if shutil.which('terminal-notifier') else 'osascript (install terminal-notifier for reliability)'}")
    print(f"python3    : {sys.executable}   <- grant THIS Accessibility permission")
    print()
    d = build_digest(cfg, day)
    for name, src in d["sources"].items():
        if src.get("available"):
            count = (src.get("count") or src.get("total_commits")
                     or src.get("session_count") or len(src.get("apps", [])) or 0)
            extra = f"source={src['source']}" if src.get("source") else ""
            print(f"  [ok]   {name:<14} {count} item(s) {extra}")
        else:
            print(f"  [--]   {name:<14} {src.get('reason', 'unavailable')}")
    h = sampler_health(cfg, day)
    if h:
        print(f"  [{'ok' if h['ok'] else '--'}]   {'sampler':<14} "
              f"{h['samples']} of ~{h['expected']} samples (ratio {h['ratio']})"
              f"{'' if h['ok'] else '  <- tracker may be down'}")
    print()
    backlog = scan_backlog(cfg, date_cls.today())
    print(f"repair backlog: {[f'{d} ({s})' for d, s, _ in backlog] or 'none'}")
    plists = list(Path("~/Library/LaunchAgents").expanduser().glob("*worklog*.plist"))
    print(f"launch agents installed: {[p.name for p in plists] or 'none'}")
    return 0


def cmd_init(cfg: dict, args) -> int:
    """Autodetect git repo locations and write them into config.json."""
    home = Path.home()
    excludes = cfg["git"].get("exclude_dir_names", DEFAULT_EXCLUDES)
    print(f"Scanning {home} for git repos (depth 4)…")
    repos = _find_repos_under(home, 4, excludes)
    if not repos:
        print("  none found — leaving git.scan_roots as-is.")
        return 0

    # use the parent dirs of discovered repos as scan roots, deduped and collapsed
    roots = sorted({str(r.parent) for r in repos})
    collapsed = [r for r in roots if not any(
        r != o and r.startswith(o.rstrip("/") + "/") for o in roots)]

    def tilde(p: str) -> str:
        return p.replace(str(home), "~", 1) if p.startswith(str(home)) else p

    print(f"  found {len(repos)} repo(s) in {len(collapsed)} location(s):")
    for r in collapsed:
        n = sum(1 for x in repos if str(x.parent) == r or str(x).startswith(r + "/"))
        print(f"    {tilde(r)}  ({n} repo(s))")

    cfg_path = expand(cfg["state_dir"]) / "config.json"
    if not cfg_path.is_file():
        cfg_path = SELF_DIR.parent / "config.json"
    with open(cfg_path) as fh:
        live = json.load(fh)
    live["git"]["scan_roots"] = [tilde(r) for r in collapsed]
    with open(cfg_path, "w") as fh:
        json.dump(live, fh, indent=2)
    print(f"  wrote git.scan_roots to {cfg_path}")

    cache = expand(cfg["state_dir"]) / "repo-cache.json"
    cache.unlink(missing_ok=True)
    return 0


def cmd_pause(cfg: dict, args) -> int:
    now = datetime.now()
    if active_pause(cfg, now):
        print(pause_status_line(cfg, now))
        print("already paused — `worklog resume` to start tracking again")
        return 0
    until = None
    if getattr(args, "until", None):
        h, m = parse_hhmm(args.until)
        until = now.replace(hour=h, minute=m, second=0, microsecond=0)
        if until <= now:
            until += timedelta(days=1)
    rec = start_pause(cfg, minutes=getattr(args, "minutes", None), until=until,
                      reason=getattr(args, "reason", "") or "")
    if not args.quiet:
        when = ("until the end of the workday" if rec["indefinite"]
                else f"until {datetime.fromisoformat(rec['until']).strftime('%H:%M')}")
        notify(cfg, "Work log paused", f"Tracking off {when}.",
               subtitle="nothing is being recorded")
    print(pause_status_line(cfg, now))
    return 0


def cmd_resume(cfg: dict, args) -> int:
    w = end_pause(cfg)
    if w is None:
        print("not paused")
        return 0
    mins = ((datetime.fromisoformat(w["ended"]) - datetime.fromisoformat(w["started"]))
            .total_seconds() / 60.0)
    if not args.quiet:
        notify(cfg, "Work log resumed", f"Tracking back on after {fmt_dur(mins)}.",
               subtitle="paused time is not recorded")
    print(f"resumed — {fmt_dur(mins)} paused")
    return 0


def pause_status_line(cfg: dict, now: datetime | None = None) -> str:
    now = now or datetime.now()
    p = active_pause(cfg, now)
    if not p:
        return "tracking"
    if p.get("indefinite") or not p.get("until"):
        return "paused indefinitely"
    left = (datetime.fromisoformat(p["until"]) - now).total_seconds() / 60.0
    return f"paused {fmt_dur(max(left, 0))} left"


def cmd_status(cfg: dict, args) -> int:
    """One compact JSON object — the menu bar renders straight from this."""
    now = datetime.now()
    day = now.date()
    p = active_pause(cfg, now)
    act = collect_activity(cfg, day)
    att = collect_meetings_attended(cfg, day)
    live = [m for m in (att.get("meetings") or [])
            if (now - datetime.fromisoformat(m["end"])).total_seconds() < 180]
    backlog = []
    try:
        backlog = [d.isoformat() for d, _, _ in scan_backlog(cfg, day)]
    except Exception:
        pass
    out = {
        "paused": bool(p),
        "indefinite": bool(p and (p.get("indefinite") or not p.get("until"))),
        "until": (p or {}).get("until"),
        "minutes_left": (None if not p or not p.get("until") else
                         max(round((datetime.fromisoformat(p["until"]) - now)
                                   .total_seconds() / 60), 0)),
        "paused_today_minutes": round(paused_minutes(
            cfg, day, *window_bounds(cfg, day)), 1),
        "state": pause_status_line(cfg, now),
        "active_today": act.get("active_duration") or "0m",
        "current_meeting": (live[-1]["title"] if live else None),
        "meetings_today": att.get("count") or 0,
        "degraded_days": backlog,
        "work_time": is_work_time(cfg, now),
        "today": day.isoformat(),
    }
    print(json.dumps(out))
    return 0


def cmd_apps(cfg: dict, args) -> int:
    """Apps seen recently, with whether each is excluded — the menu renders this."""
    acfg = cfg["activity"]
    days = int(getattr(args, "days", 7) or 7)
    today = date_cls.today()
    seen: dict = {}
    for i in range(days):
        day = today - timedelta(days=i)
        path = expand(cfg["state_dir"]) / "raw" / day.isoformat() / "activity.jsonl"
        if not path.is_file():
            continue
        for line in path.read_text(errors="replace").splitlines():
            if not line.strip():
                continue
            try:
                r = json.loads(line)
            except json.JSONDecodeError:
                continue
            app = r.get("app") or ""
            if not app or app.startswith("("):
                continue
            e = seen.setdefault(app, {"app": app, "bundle": r.get("bundle", ""),
                                      "samples": 0})
            e["samples"] += 1
            if r.get("bundle"):
                e["bundle"] = r["bundle"]
    # anything currently excluded should still be listed, so it can be switched back on
    for name in acfg.get("exclude_apps") or []:
        seen.setdefault(str(name), {"app": str(name), "bundle": "", "samples": 0})
    out = []
    for e in seen.values():
        e["excluded"] = is_excluded_app(acfg, e["app"], e.get("bundle"))
        e["minutes"] = e["samples"]          # one sample per minute
        out.append(e)
    out.sort(key=lambda e: (-e["samples"], e["app"].lower()))
    print(json.dumps(out))
    return 0


def _write_config(cfg: dict) -> Path:
    """Persist the live config. This file is untracked; the repo ships an example."""
    for cand in (DEFAULT_STATE / "config.json", SELF_DIR.parent / "config.json"):
        if cand.is_file():
            cand.write_text(json.dumps(cfg, indent=2) + "\n", encoding="utf-8")
            return cand
    raise SystemExit("worklog: no config.json to write")


def cmd_config(cfg: dict, args) -> int:
    action = args.action
    acfg = cfg["activity"]
    if action == "get":
        node = cfg
        if args.key:
            for part in args.key.split("."):
                if not isinstance(node, dict) or part not in node:
                    die(f"no such key: {args.key}")
                node = node[part]
        print(json.dumps(node, indent=2))
        return 0

    if action in ("exclude-app", "include-app"):
        # `config exclude-app Notes` puts the name in the first positional, while
        # `config set k v` uses both, so accept it from either slot
        name = ((args.key or "") + " " + (args.value or "")).strip()
        if not name:
            die("give an app name")
        lst = acfg.setdefault("exclude_apps", [])
        low = [str(x).lower() for x in lst]
        allow = acfg.setdefault("vendor_allow_bundles", [])
        bid = bundle_for_app(name) or ""
        if action == "exclude-app":
            if name.lower() not in low:
                lst.append(name)
            # an Apple app is kept by the allow list, so excluding it means
            # removing it from there as well or the vendor rule would still pass it
            if bid:
                acfg["vendor_allow_bundles"] = [a for a in allow
                                                if str(a).lower() != bid.lower()]
        else:
            acfg["exclude_apps"] = [x for x in lst if str(x).lower() != name.lower()]
            prefixes = acfg.get("vendor_exclude_prefixes") or []
            if bid and any(bid.lower().startswith(str(p).lower()) for p in prefixes):
                if bid not in acfg.setdefault("vendor_allow_bundles", []):
                    acfg["vendor_allow_bundles"].append(bid)
        where = _write_config(cfg)
        state = "excluded" if is_excluded_app(acfg, name, bid) else "included"
        log_line(cfg, f"config: {name} is now {state}")
        print(f"{name} -> {state} ({where.name})")
        return 0

    if action == "set":
        if not args.key:
            die("give a dotted key, e.g. work_hours.end")
        try:
            val = json.loads(args.value)
        except (json.JSONDecodeError, TypeError):
            val = args.value
        node = cfg
        parts = args.key.split(".")
        for part in parts[:-1]:
            node = node.setdefault(part, {})
            if not isinstance(node, dict):
                die(f"{args.key} is not a nested key")
        node[parts[-1]] = val
        where = _write_config(cfg)
        log_line(cfg, f"config: set {args.key} = {val!r}")
        print(f"{args.key} = {json.dumps(val)} ({where.name})")
        return 0
    die(f"unknown config action {action}")


def cmd_notify_test(cfg: dict, args) -> int:
    backend = "terminal-notifier" if shutil.which("terminal-notifier") else "osascript"
    print(f"Sending test notification via {backend}…")
    notify(cfg, "Work log", "Notifications are working.",
           subtitle="test", open_path=str(expand(cfg["output_dir"])))
    print("If nothing appeared, check System Settings → Notifications and allow "
          f"{'terminal-notifier' if backend == 'terminal-notifier' else 'Script Editor'}.")
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(prog="worklog", description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)

    p = sub.add_parser("sample", help="take one activity sample")
    p.add_argument("--force", action="store_true", help="ignore work-hours window")

    for name, helptext in (("report", "write the markdown work log"),
                           ("digest", "print raw JSON digest"),
                           ("weekly", "write the week's merged MIS table"),
                           ("doctor", "diagnose sources and permissions")):
        q = sub.add_parser(name, help=helptext)
        q.add_argument("--date", default="today",
                       help="YYYY-MM-DD | today | yesterday")
        if name == "report":
            q.add_argument("--quiet", action="store_true",
                           help="write the log but send no notification")

    q = sub.add_parser("repair", help="redo days whose log is missing or unsummarized")
    q.add_argument("--date", default=None,
                   help="YYYY-MM-DD (default: scan the lookback window)")
    q.add_argument("--dry-run", action="store_true",
                   help="list what would be repaired, change nothing")
    q.add_argument("--quiet", action="store_true",
                   help="repair but send no notification")

    q = sub.add_parser("pause", help="stop recording during personal time")
    g = q.add_mutually_exclusive_group()
    g.add_argument("--minutes", type=float, help="pause for this many minutes")
    g.add_argument("--until", help="pause until HH:MM")
    g.add_argument("--indefinite", action="store_true",
                   help="pause until resumed (still ends at the workday's end)")
    q.add_argument("--reason", default="", help="note stored with the pause window")
    q.add_argument("--quiet", action="store_true", help="send no notification")

    q = sub.add_parser("resume", help="start recording again")
    q.add_argument("--quiet", action="store_true", help="send no notification")

    sub.add_parser("status", help="print current state as JSON (for the menu bar)")

    q = sub.add_parser("apps", help="apps seen recently and whether each is excluded")
    q.add_argument("--days", type=int, default=7)

    q = sub.add_parser("config", help="read or change settings")
    q.add_argument("action", choices=["get", "set", "exclude-app", "include-app"])
    q.add_argument("key", nargs="?", default="")
    q.add_argument("value", nargs="?", default="")
    sub.add_parser("init", help="autodetect git repo locations into config")
    sub.add_parser("notify-test", help="send a test notification")

    args = ap.parse_args()
    cfg = load_config()
    return {"sample": cmd_sample, "report": cmd_report, "digest": cmd_digest,
            "doctor": cmd_doctor, "repair": cmd_repair, "weekly": cmd_weekly,
            "pause": cmd_pause, "resume": cmd_resume, "status": cmd_status,
            "apps": cmd_apps, "config": cmd_config,
            "init": cmd_init, "notify-test": cmd_notify_test}[args.cmd](cfg, args)


if __name__ == "__main__":
    sys.exit(main())
